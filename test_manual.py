from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
prs.save("manual_test.pptx")
