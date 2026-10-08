#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inventory the template's 23 slides and emit the exemplar manifest.

Run this BEFORE designing slot names. layout_manifest.json records only shape *counts*
for the slides; the design lives in the shapes themselves and nobody has looked at them.

    python -m aippt.analyze_exemplars --dump          # read this. then curate.
    python -m aippt.analyze_exemplars                 # write/merge the manifest

Read-only by construction: the template is slurped to BytesIO, never opened by path.

The manifest it writes is UNCURATED -- every shape lands in `slots` with `slot_id: null`.
A human then assigns semantic slot_ids. Re-running merges on `shape_id` and NEVER clobbers
a hand-assigned slot_id (that is the whole reason this is a merge and not a rewrite).
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
from pathlib import Path
from typing import Any

from pptx.oxml.ns import qn

from .analyze_pptx_master import EMU_PER_INCH, content_budget, emu_to_in, find_default_template

#: Slides with 41-51 shapes are bespoke hand-drawn figures (brief §1). Cloning one and
#: hoping to swap the text produces garbage, and each is 40+ shapes of mapping ambiguity
#: for zero reuse value.
MAX_EXEMPLAR_SHAPES = 12

#: The nine roles the HTML adapter needs to render the 30-card corpus (brief §3.3).
REQUIRED_ROLES = (
    "deck-title", "section-divider", "content-bullets", "content-two-col", "content-table",
    "content-image", "content-steps", "content-cards", "content-code",
)


class ExemplarRejected(ValueError):
    """A slide cannot be an exemplar. The reason is the message."""


def _effective_style(shape: Any) -> dict[str, Any] | None:
    """The run style read off the SLIDE's own <a:rPr> -- this is where the design is.

    Not the theme, not the master: Q3 F-03 established that the brand red and the
    뷰웍스 B / 뷰웍스 M typefaces live in slide- and layout-level run properties. The theme's
    accent slots are stock Office blue.
    """
    if not getattr(shape, "has_text_frame", False):
        return None
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            rPr = run._r.find(qn("a:rPr"))
            if rPr is None:
                continue
            latin = rPr.find(qn("a:latin"))
            srgb = rPr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
            sz = rPr.get("sz")
            return {
                "font_pt": (int(sz) / 100.0) if sz else None,
                "typeface": latin.get("typeface") if latin is not None else None,
                "color": (srgb.get("val") or "").upper() or None if srgb is not None else None,
                "bold": rPr.get("b") == "1",
                "source": "slide-run-rPr",
            }
    return None


def _shape_kind(shape: Any) -> str:
    if getattr(shape, "has_table", False):
        return "table"
    if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
        return "picture"
    if shape._element.tag == qn("p:grpSp"):
        return "group"
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "other"


def _reject_unclonable(slide: Any, index: int) -> None:
    """Charts and OLE objects are separate PARTS, and clone.py shares parts between the
    source and the clone (brief §2.2). Two clones of a chart-bearing exemplar would share
    one chart part, so rebranding one would recolour the other. Refuse the exemplar rather
    than build a chart path we cannot test."""
    for tag, what in ((qn("c:chart"), "a chart"), (qn("p:oleObj"), "an OLE object")):
        if slide._element.find(f".//{tag}") is not None:
            raise ExemplarRejected(
                f"slide {index} contains {what}. Charts and OLE objects live in their own "
                f"part; cloning shares the part between source and clone, so mutating one "
                f"clone would mutate the other. Phase 1 exemplars may not contain them."
            )


def inventory_slide(slide: Any, index: int) -> dict[str, Any]:
    _reject_unclonable(slide, index)

    slots: list[dict[str, Any]] = []
    seen_ids: set[int] = set()
    for z, shape in enumerate(slide.shapes):
        sid = shape.shape_id
        if sid in seen_ids:
            # PowerPoint can emit duplicate cNvPr ids across a group boundary. shape_id is
            # our primary key, so a duplicate is not something to work around.
            raise ExemplarRejected(
                f"slide {index} has two top-level shapes with cNvPr id {sid}. Re-save the "
                f"slide in PowerPoint to renumber them."
            )
        seen_ids.add(sid)

        emu = {"left": shape.left, "top": shape.top,
               "width": shape.width, "height": shape.height}
        emu = {k: (None if v is None else int(v)) for k, v in emu.items()}
        style = _effective_style(shape)
        kind = _shape_kind(kind_shape := shape)
        text = kind_shape.text_frame.text if getattr(shape, "has_text_frame", False) else None

        slots.append({
            "slot_id": None,                       # <-- the human fills this in
            "role": None,                          # or "decoration"
            "kind": kind,
            "shape_id": sid,
            "shape_name": shape.name,              # fingerprint only. NOT unique. Never a key.
            "z_index": z,                          # fingerprint only
            "is_placeholder": shape.is_placeholder,
            "placeholder_idx": (shape.placeholder_format.idx
                                if shape.is_placeholder else None),
            "emu": emu,
            "inches": {k: emu_to_in(v) for k, v in emu.items()},
            "effective_style": style,
            "budget": (content_budget(emu["width"], emu["height"],
                                      (style or {}).get("font_pt") or 18.0, "BODY")
                       if kind == "text" else None),
            "sample_text": text,                   # the leftover-text test (V-4) asserts on this
            "required": False,
            "on_unbound": "remove",
        })

    return {
        "exemplar_id": None,                       # <-- the human names this
        "source_slide_index": index,
        "layout_name": slide.slide_layout.name,
        "roles": [],
        "supports_blocks": [],
        "shape_count": len(slots),
        "slots": slots,
        "brand_shapes": [],
        "decorations": [],
    }


def analyze(template: Path) -> dict[str, Any]:
    from pptx import Presentation

    raw = template.read_bytes()
    prs = Presentation(io.BytesIO(raw))          # BytesIO => the file on disk is never touched

    exemplars, rejected = [], []
    for i, slide in enumerate(prs.slides):
        if len(slide.shapes) > MAX_EXEMPLAR_SHAPES:
            rejected.append({"source_slide_index": i, "reason":
                             f"{len(slide.shapes)} top-level shapes -- a bespoke figure, "
                             f"not a reusable exemplar (brief §1)"})
            continue
        try:
            exemplars.append(inventory_slide(slide, i))
        except ExemplarRejected as exc:
            rejected.append({"source_slide_index": i, "reason": str(exc)})

    return {
        "schema_version": "2.0",
        "generated_by": "aippt.analyze_exemplars",
        "curated": False,                          # slots.py REFUSES to load until this is true
        "source": {"filename": template.name,
                   "sha256": hashlib.sha256(raw).hexdigest(),
                   "bytes": len(raw)},
        "required_roles": list(REQUIRED_ROLES),
        "exemplars": exemplars,
        "rejected_slides": rejected,
    }


def merge(new: dict[str, Any], old: dict[str, Any]) -> dict[str, Any]:
    """Re-run against the same template WITHOUT destroying an hour of curation.

    Merge key is (source_slide_index, shape_id). Hand-assigned fields win; everything the
    analyzer measures is refreshed. If a shape_id has vanished, the slot is dropped and the
    loss is reported -- silently keeping a slot that no longer resolves is how you get a
    TemplateDriftError three modules downstream instead of here.
    """
    hand = ("exemplar_id", "roles", "supports_blocks", "brand_shapes", "decorations")
    hand_slot = ("slot_id", "role", "required", "on_unbound")

    old_ex = {e["source_slide_index"]: e for e in old.get("exemplars", [])}
    lost: list[str] = []

    for ex in new["exemplars"]:
        prev = old_ex.get(ex["source_slide_index"])
        if prev is None:
            continue
        for key in hand:
            if prev.get(key):
                ex[key] = prev[key]
        prev_slots = {s["shape_id"]: s for s in prev.get("slots", [])}
        for slot in ex["slots"]:
            p = prev_slots.pop(slot["shape_id"], None)
            if p:
                for key in hand_slot:
                    if p.get(key) is not None:
                        slot[key] = p[key]
        for orphan in prev_slots.values():
            if orphan.get("slot_id"):
                lost.append(f"slide {ex['source_slide_index']}: slot "
                            f"{orphan['slot_id']!r} (shape_id {orphan['shape_id']}) no longer "
                            f"exists in the template")

    new["curated"] = old.get("curated", False)
    if lost:
        new["curated"] = False                     # a lost slot invalidates the curation
        new["merge_losses"] = lost
    return new


def render_catalog(manifest: dict[str, Any]) -> str:
    out = ["# Exemplar Catalog", "",
           f"Generated from `{manifest['source']['filename']}` "
           f"(sha256 `{manifest['source']['sha256'][:12]}…`) by `analyze_exemplars.py`.",
           f"**Curated: {manifest['curated']}** — the renderer refuses to run until this is true.",
           ""]
    for ex in manifest["exemplars"]:
        out += [f"## slide {ex['source_slide_index']} — `{ex['exemplar_id'] or 'UNNAMED'}` "
                f"({ex['shape_count']} shapes, layout `{ex['layout_name']}`)", "",
                "| slot_id | kind | shape_id | name | z | box (in) | style | sample text |",
                "|---|---|---|---|---|---|---|---|"]
        for s in ex["slots"]:
            box = s["inches"]
            st = s["effective_style"] or {}
            style = (f"{st.get('font_pt')}pt {st.get('typeface') or ''} "
                     f"#{st.get('color') or '—'}") if st else "—"
            sample = (s["sample_text"] or "").replace("\n", " ⏎ ")[:40]
            out.append(f"| {s['slot_id'] or '**?**'} | {s['kind']} | {s['shape_id']} | "
                       f"`{s['shape_name']}` | {s['z_index']} | "
                       f"{box['left']},{box['top']} {box['width']}×{box['height']} | "
                       f"{style} | {sample} |")
        out.append("")
    if manifest["rejected_slides"]:
        out += ["## Rejected slides", ""]
        out += [f"- **{r['source_slide_index']}** — {r['reason']}"
                for r in manifest["rejected_slides"]]
        out.append("")
    return "\n".join(out) + "\n"


def print_dump(manifest: dict[str, Any]) -> None:
    """§8 step 1. READ THIS BEFORE YOU NAME A SINGLE SLOT."""
    for ex in manifest["exemplars"]:
        print(f"\n{'=' * 78}\nSLIDE {ex['source_slide_index']}  "
              f"layout={ex['layout_name']!r}  shapes={ex['shape_count']}\n{'=' * 78}")
        for s in ex["slots"]:
            st = s["effective_style"] or {}
            print(f"  [z={s['z_index']}] id={s['shape_id']:<4} {s['kind']:<8} "
                  f"{s['shape_name']!r}")
            print(f"        box   : {s['inches']}")
            if st:
                print(f"        style : {st.get('font_pt')}pt {st.get('typeface')!r} "
                      f"#{st.get('color')} bold={st.get('bold')}")
            if s["sample_text"]:
                print(f"        text  : {s['sample_text'][:70]!r}")
    print(f"\n{len(manifest['exemplars'])} candidate exemplars, "
          f"{len(manifest['rejected_slides'])} rejected.")
    print(f"Roles still to cover: {', '.join(REQUIRED_ROLES)}")


def main(argv: list[str] | None = None) -> int:
    repo_root = Path(__file__).resolve().parents[2]
    ap = argparse.ArgumentParser(description="Inventory the template's slides as clone exemplars.")
    ap.add_argument("--template", type=Path, default=None)
    ap.add_argument("--out-dir", type=Path, default=repo_root / "reference")
    ap.add_argument("--dump", action="store_true", help="print the inventory and write nothing")
    args = ap.parse_args(argv)

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # Hangul on a Windows console
    except (AttributeError, ValueError):
        pass

    template = args.template or find_default_template(repo_root)
    if template is None or not Path(template).is_file():
        print("ERROR: template not found. Pass --template <path>.", file=sys.stderr)
        return 2

    manifest = analyze(Path(template).resolve())

    if args.dump:
        print_dump(manifest)
        return 0

    json_path = args.out_dir / "exemplar_manifest.json"
    if json_path.is_file():
        manifest = merge(manifest, json.loads(json_path.read_text(encoding="utf-8")))
        for loss in manifest.get("merge_losses", []):
            print(f"CURATION LOST: {loss}", file=sys.stderr)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    (args.out_dir / "exemplar_catalog.md").write_text(render_catalog(manifest), encoding="utf-8")
    print(f"wrote {json_path}\nwrote {args.out_dir / 'exemplar_catalog.md'}")
    if not manifest["curated"]:
        print("\nNOT CURATED. Assign slot_ids + exemplar_ids, then set \"curated\": true.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())