#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""One-off build step. Produces reference/renderer_template.pptx (0 slides).

    python src/aippt/tools/make_renderer_template.py --write-pin

Run once, commit the output, pin its SHA in reference/template_pin.json. NEVER import this from
the request path: it is the ONLY sanctioned use of python-pptx's private slide-deletion surface
(critic rule 3). The source template is opened read-only and never written back -- the save
target is asserted to be a different path before anything is saved.

Why a separate artifact exists at all: Presentation(SRC) opens a deck that ALREADY HAS 23
exemplar slides, and the public Slides API has no delete operation. add_slide() would append a
24th, every export would ship the other 23, and the inherited page-number field on the new slide
would read 24.

Why the pin is written from the observed package and never from a literal typed in by hand: a
hand-copied signature pins what someone believed, and the whole point of the pin is to catch the
day that belief stops being true. Everything under "layout_signature" below is read back off the
package this script just built.

Rebuilds are NOT byte-reproducible -- ZIP entries carry timestamps -- so we pin the committed
artifact, not a rebuild. Re-pinning is therefore a deliberate, reviewed act (plan Sec 4.3).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

# <root>/src/aippt/tools/make_renderer_template.py -> <root>. Anchored to the file, not to cwd,
# so the builder produces the same artifact from any working directory.
PROJECT_ROOT = Path(__file__).resolve().parents[3]

SRC = PROJECT_ROOT / "Vieworks_R&D_ppt_Template.pptx"
DST = PROJECT_ROOT / "reference" / "renderer_template.pptx"
PIN = PROJECT_ROOT / "reference" / "template_pin.json"

SRC_SHA = "0de7d70cf45197e959d93c31eedd7599c42c23c1153d7b67b18d332141e61442"
SRC_BYTES = 1_881_082

BUILDER = "aippt.tools.make_renderer_template@1"
PIN_VERSION = 1

# The layout the renderer draws on: "2_사용자 지정 레이아웃", the one 15 of the 23 template
# slides already use. Its chrome (gradient divider, SVG logo, page number) is what we inherit.
LAYOUT_INDEX = 3
TITLE_PH_IDX = 10

# Sanity anchors, NOT the pinned values. The pinned values are observed (see observe_signature).
# These only assert that layout 3 is still the layout we think it is -- if the template is re-cut
# and layout 3 becomes something else, we want a loud failure here rather than a confident pin of
# the wrong layout.
EXPECT_LAYOUT_PARTNAME = "/ppt/slideLayouts/slideLayout4.xml"
EXPECT_MEDIA = ("image5.jpeg", "image2.svg")  # divider JPEG, logo SVG

_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class BuildError(RuntimeError):
    """The build refuses to produce an artifact it cannot vouch for."""


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def delete_all_slides(prs: Presentation) -> int:
    """Remove every slide from an open Presentation. Returns the count removed.

    THE LANDMINE. `Presentation.slides` is not a plain accessor. Every evaluation of it runs

        self.part.rename_slide_parts([sldId.rId for sldId in sldIdLst])

    which does a `related_part(rId)` lookup for every <p:sldId> still in the list. Evaluate
    `prs.slides` at any moment when a relationship has been dropped but its <p:sldId> is still
    present and that lookup raises KeyError on a package that is only transiently inconsistent.
    (In python-pptx 1.0.2 `slides` happens to be a @lazyproperty, so the cache hides this today.
    That is an implementation detail of one version, not a contract, and it is not what this code
    relies on.)

    So: bind the XML element ONCE, mutate only the element, and do not read `prs.slides` again
    until the loop has finished and the package is consistent.

    Why drop_rel AND the sldId removal, rather than just the sldId: python-pptx serializes every
    part reachable through the relationship graph. Removing only the <p:sldId> entry leaves the
    presentation.xml.rels relationship intact, so all 23 slide parts would still be written into
    the package -- invisible in PowerPoint, but dead weight and a source of confusing audits.
    """
    sld_id_lst = prs.slides._sldIdLst  # bind once; see docstring
    removed = 0
    for sld_id in list(sld_id_lst):  # list(): never mutate the tree you are iterating
        # drop_rel first, while the <p:sldId> still carries the r:id. XmlPart.drop_rel pops the
        # rel only when its reference count is < 2, and right now it is exactly 1.
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)
        removed += 1
    return removed


def _layout_media(layout: Any) -> list[str]:
    """Media filenames the layout part relates to, in sorted order.

    Note this is normally THREE files, not two: the logo is an SVG `p:pic`, and an SVG in OOXML
    carries a raster fallback alongside it (the `svgBlip` extension). Both the .svg and its
    fallback are related to the layout, and both matter -- which one a converter honours is
    exactly the open question in plan Sec 7.1 assertion 5.
    """
    return sorted(
        posixpath.basename(str(rel.target_partname))
        for rel in layout.part.rels.values()
        if not rel.is_external and "/media/" in str(rel.target_partname)
    )


def _accent1(prs: Presentation) -> str | None:
    """Theme accent1. Informational in the pin: it is the colour that must NEVER appear in a
    generated shape, because add_shape() hands it to you by default (plan Sec 6.4)."""
    try:
        theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
        val = etree.fromstring(theme.blob).xpath(
            "//a:clrScheme/a:accent1/a:srgbClr/@val", namespaces=_NS
        )
        return str(val[0]).upper() if val else None
    except Exception:  # noqa: BLE001 -- a cosmetic note must not fail the build
        return None


def observe_signature(prs: Presentation) -> dict[str, Any]:
    """Read the layout signature back off the package. This -- not a hand-typed literal -- is
    what gets pinned."""
    layout = prs.slide_layouts[LAYOUT_INDEX]

    ph = layout.placeholders.get(TITLE_PH_IDX)
    # .get(idx), NOT placeholders[10]. On a LAYOUT, `placeholders[...]` is positional (it comes
    # from _BaseShapes.__getitem__) and layout 3 has exactly one placeholder, so [10] raises
    # IndexError. Only _SlidePlaceholders.__getitem__ is keyed by idx. Easy to get backwards.
    if ph is None:
        raise BuildError(
            f"layout {LAYOUT_INDEX} has no placeholder with idx {TITLE_PH_IDX}; "
            "the title has nowhere to go and the renderer cannot work against this template"
        )

    return {
        "layout_index": LAYOUT_INDEX,
        "layout_name": layout.name,
        "layout_partname": str(layout.part.partname),
        "slide_width_emu": int(prs.slide_width),
        "slide_height_emu": int(prs.slide_height),
        "slide_count": len(prs.slides),
        "layout_count": len(prs.slide_layouts),
        "placeholder_idxs": sorted(p.placeholder_format.idx for p in layout.placeholders),
        f"placeholder_{TITLE_PH_IDX}_emu": {
            "left": int(ph.left),
            "top": int(ph.top),
            "width": int(ph.width),
            "height": int(ph.height),
        },
        "required_media": _layout_media(layout),
    }


def check_signature(sig: dict[str, Any]) -> None:
    """Refuse to pin a package that is not the one this builder was written against."""
    problems: list[str] = []
    if sig["slide_count"] != 0:
        problems.append(f"slides survived deletion: {sig['slide_count']}")
    if sig["layout_count"] != 5:
        problems.append(f"layout count is {sig['layout_count']}, expected 5")
    if (sig["slide_width_emu"], sig["slide_height_emu"]) != (12192000, 6858000):
        problems.append("slide size is not 12192000 x 6858000 EMU (16:9, 960 x 540 pt)")
    if sig["layout_partname"] != EXPECT_LAYOUT_PARTNAME:
        problems.append(
            f"layout {LAYOUT_INDEX} is {sig['layout_partname']}, expected {EXPECT_LAYOUT_PARTNAME}"
        )
    missing = [m for m in EXPECT_MEDIA if m not in sig["required_media"]]
    if missing:
        problems.append(f"layout {LAYOUT_INDEX} lost media {missing}; has {sig['required_media']}")
    if problems:
        raise BuildError("refusing to pin:\n  - " + "\n  - ".join(problems))


def audit_package(path: Path) -> None:
    """Acceptance gate (plan Sec 2.3), checks 1-4, run against the saved ZIP.

    Checks 1-3 are re-read through python-pptx by the caller; this function is the part that has
    to look at the raw package, because "no orphaned parts" and "no dangling relationships" are
    properties of the ZIP, not of the object model.
    """
    problems: list[str] = []
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())

        strays = sorted(n for n in names if n.startswith(("ppt/slides/", "ppt/notesSlides/")))
        if strays:
            problems.append(f"slide/notes parts survived into the package: {strays}")

        for media in EXPECT_MEDIA:
            if f"ppt/media/{media}" not in names:
                problems.append(f"ppt/media/{media} is missing from the package")

        # Dangling relationships: every internal Target must resolve to a part that exists.
        for rels_name in sorted(n for n in names if n.endswith(".rels")):
            part_dir = posixpath.dirname(posixpath.dirname(rels_name))  # ppt/_rels/x.rels -> ppt
            for rel in etree.fromstring(zf.read(rels_name)):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target") or ""
                resolved = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join(part_dir, target))
                )
                if resolved not in names:
                    problems.append(f"{rels_name}: dangling rel -> {target}")

    if problems:
        raise BuildError("package audit failed:\n  - " + "\n  - ".join(problems))


def has_slide_ph(slide: Any, idx: int) -> bool:
    """True if the slide carries a placeholder with this idx.

    The two placeholder collections are exact opposites and there is no way to write this once:
      - slide.placeholders[idx]  -> keyed by IDX, raises KeyError, and has NO .get()
      - layout.placeholders[idx] -> keyed by POSITION (from _BaseShapes), and .get(idx) is the
                                    only idx lookup it offers
    Mixing them up gives you an IndexError from a line that reads like it works.
    """
    try:
        slide.placeholders[idx]
    except KeyError:
        return False
    return True


def smoke_add_slide(path: Path) -> None:
    """Acceptance gate check 6: add_slide(layout 3) -> save -> reopen -> it is slide 1.

    The page-number FIELD is rendered by PowerPoint, not by python-pptx, so we cannot assert the
    glyph "1" here -- what we can assert is the only thing that determines it: the new slide is
    the one and only slide, at index 0. Reading the rendered "1" is a human check
    (HUMAN_VERIFICATION.md).
    """
    prs = Presentation(str(path))
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    if not has_slide_ph(slide, TITLE_PH_IDX):
        raise BuildError(f"add_slide did not clone placeholder idx {TITLE_PH_IDX} onto the slide")

    with tempfile.TemporaryDirectory() as tmp:
        probe = Path(tmp) / "smoke.pptx"
        prs.save(str(probe))
        out = Presentation(str(probe))
        problems: list[str] = []
        if len(out.slides) != 1:
            problems.append(f"expected exactly 1 slide, got {len(out.slides)}")
        else:
            s = out.slides[0]
            if str(s.slide_layout.part.partname) != EXPECT_LAYOUT_PARTNAME:
                problems.append(f"slide 1 relates to {s.slide_layout.part.partname}")
            if not has_slide_ph(s, TITLE_PH_IDX):
                problems.append(f"placeholder idx {TITLE_PH_IDX} did not survive the round trip")
        if problems:
            raise BuildError("smoke test failed:\n  - " + "\n  - ".join(problems))


def build() -> dict[str, Any]:
    if DST.resolve() == SRC.resolve():
        raise BuildError("destination is the source template; refusing to overwrite it")

    if not SRC.is_file():
        raise BuildError(f"source template not found: {SRC}")

    lock = SRC.with_name("~$" + SRC.name)
    if lock.exists():
        # An Office lock file. Never hashed, never read -- but its presence means the template is
        # probably open in PowerPoint, which is worth saying out loud before we read it.
        print(f"warning: {lock.name} exists -- close PowerPoint before trusting this build.",
              file=sys.stderr)

    src_sha = sha256_file(SRC)
    if src_sha != SRC_SHA:
        raise BuildError(
            f"source template changed\n  expected {SRC_SHA}\n  actual   {src_sha}\n"
            "This builder was written against the expected package. Re-audit the template "
            "(aippt.analyze_pptx_master), update SRC_SHA deliberately, then rebuild."
        )

    prs = Presentation(str(SRC))  # SRC is opened read-only and never saved back
    removed = delete_all_slides(prs)

    DST.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(DST))

    out = Presentation(str(DST))  # re-read: audit the artifact, not the in-memory object
    sig = observe_signature(out)
    check_signature(sig)
    audit_package(DST)
    smoke_add_slide(DST)

    return {
        "slides_removed": removed,
        "source_template": {
            "filename": SRC.name,
            "sha256": src_sha,
            "bytes": SRC.stat().st_size,
        },
        "renderer_template": {
            "filename": DST.relative_to(PROJECT_ROOT).as_posix(),
            "sha256": sha256_file(DST),
            "bytes": DST.stat().st_size,
            "built_from_sha256": src_sha,
            "builder": BUILDER,
        },
        "layout_signature": sig,
        "theme": {
            "accent1": _accent1(out),
            "note": "Office blue. MUST NOT appear in any generated shape.",
        },
    }


def write_pin(result: dict[str, Any]) -> None:
    pin = {
        "pin_version": PIN_VERSION,
        "source_template": result["source_template"],
        "renderer_template": result["renderer_template"],
        "layout_signature": result["layout_signature"],
        "theme": result["theme"],
    }
    PIN.parent.mkdir(parents=True, exist_ok=True)
    PIN.write_text(json.dumps(pin, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build reference/renderer_template.pptx (0 slides).")
    ap.add_argument(
        "--write-pin",
        action="store_true",
        help="also write reference/template_pin.json from the OBSERVED package",
    )
    args = ap.parse_args(argv)

    try:
        result = build()
    except BuildError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    # Always emit the observed signature, so the pin can be reviewed before it is written.
    print(json.dumps(result, indent=2, ensure_ascii=False))

    if args.write_pin:
        write_pin(result)
        print(f"\nwrote {PIN.relative_to(PROJECT_ROOT).as_posix()}", file=sys.stderr)
    else:
        print(
            "\nnot written: re-run with --write-pin to update reference/template_pin.json",
            file=sys.stderr,
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
