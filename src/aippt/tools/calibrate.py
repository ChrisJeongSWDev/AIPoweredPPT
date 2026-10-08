#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The calibration slide (plan Sec 7.2). Phase 0, Step 2 -- blocking, before any feature work.

    python src/aippt/tools/calibrate.py            # -> reference/calibration.pptx + .json

This converts "I think the coordinate mapping is right" into a measured number. It emits a
fixture deck whose every shape has a known, exact position in the 960 x 540 logical space, plus a
JSON sidecar giving the ground truth. Export it through the ACTUAL target consumer (PowerPoint,
and the LibreOffice path if that is what ships), raster it to PNG, and diff against a screenshot
of the HTML stage at zoom 1. That diff is the mapping's error bar. Acceptance is <= 2 px
(Sec 7.5) -- and coordinate equality alone does not count as evidence (critic rule 12).

What each fixture probes:
  corners     the four exact slide corners -- catches an off-by-one in the model->EMU origin
  strokes     widths 0, 1, 4, 8 -- catches box-model drift (does the stroke grow the outer edge?)
  rotation    30 deg and 45 deg -- catches a centre-vs-corner rotation-origin mismatch
  wrap        a known string in a known box at 14 pt -- catches a text-metric mismatch, which
              shows up as a wrapped line count that differs between browser and PowerPoint
  mixed       a KO/EN string -- catches a missing latin/ea font split (the Hangul would fall back
              to a different face, and the line width would move with it)
  chrome      the divider and logo, inherited from layout 3 for free -- undistorted or bust

PHASE 0 ISOLATION. canvas/geometry.py and canvas/style.py are Phase 1 and do not exist yet, so
the EMU conversion and the four add_shape overrides are INLINED below. When Phase 1 lands, these
move out and this file imports them; until then, Phase 0 stands alone and this file is the only
place that knows them. Keep the duplication obvious rather than inventing the Phase 1 modules
early -- the point is that Phase 0's gates pass on their own.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

if __package__ in (None, ""):  # `python src/aippt/tools/calibrate.py`, no install
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from aippt.canvas.pin import (  # noqa: E402
    TITLE_PH_IDX,
    TemplatePinError,
    load_pin,
    open_renderer_template,
)

PROJECT_ROOT = Path(__file__).resolve().parents[3]
DST = PROJECT_ROOT / "reference" / "calibration.pptx"

# --- inlined from what will become canvas/geometry.py (Phase 1, task 1.2) --------------------
# 1 logical unit == 1 pt == exactly 12,700 EMU. 12192000/960 == 6858000/540 == 914400/72.
# Integer and exact by construction, which is why the round-trip tolerance below is 0 EMU and
# not "close enough".
EMU_PER_UNIT = 12700
CANVAS_W = 960
CANVAS_H = 540


def emu(u: float) -> Emu:
    return Emu(round(u * EMU_PER_UNIT))


# --- inlined from what will become canvas/style.py (Phase 1, task 1.3) -----------------------
LATIN = "Century Gothic"
EA = "맑은 고딕"  # Century Gothic ships NO Hangul glyphs. This is not optional.

BRAND_RED = "FF0032"
INK = "000000"
FORBIDDEN = "4472C4"  # theme accent1. If this appears in the output, an override was missed.


def set_run_fonts(run: Any, latin: str = LATIN, ea: str = EA) -> None:
    """The latin/EA split. python-pptx cannot express an East-Asian typeface -- Font.name writes
    only <a:latin>, and CT_TextCharacterProperties declares no `ea` child at all -- so <a:ea> has
    to be inserted by hand.

    The successor list is the DrawingML child order for CT_TextCharacterProperties: <a:ea> must
    follow <a:latin> and precede <a:cs>/<a:sym>/<a:hlink*>/<a:rtl>/<a:extLst>. Get it wrong and
    PowerPoint calls the file corrupt. (insert_element_before appends when it finds none of the
    successors, so the list must be complete, not merely plausible.)
    """
    rPr = run._r.get_or_add_rPr()
    rPr.get_or_add_latin().typeface = latin
    ea_el = rPr.find(qn("a:ea"))
    if ea_el is None:
        ea_el = rPr.makeelement(qn("a:ea"), {})
        rPr.insert_element_before(
            ea_el, "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"
        )
    ea_el.set("typeface", ea)


def style_text_frame(tf: Any, text: str | None, size_pt: float = 14) -> None:
    """The margin trap: PowerPoint defaults to 0.1 in left/right and 0.05 in top/bottom, which is
    7.2 and 3.6 logical units. A browser div with padding:0 therefore wraps at a DIFFERENT width
    than PowerPoint does, and the result is an off-by-one-line overflow that presents as a
    coordinate bug. Set all four explicitly, and mirror them in the CSS."""
    tf.margin_left = tf.margin_right = Pt(7.2)
    tf.margin_top = tf.margin_bottom = Pt(3.6)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    if text is None:
        return
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(0)  # explicit -- do not inherit
    p.space_after = Pt(0)
    p.line_spacing = 1.2  # == CSS line-height: 1.2
    run = p.add_run()
    run.text = text
    set_run_fonts(run)
    run.font.size = Pt(size_pt)
    run.font.bold = False
    run.font.color.rgb = RGBColor.from_string(INK)


def draw_rect(slide: Any, el: dict[str, Any]) -> Any:
    """add_shape() hands you Office blue #4472C4, blue-outlined, drop-shadowed, with white
    theme-font text -- because new_autoshape_sp hard-codes a <p:style> block pointing
    lnRef/fillRef/effectRef at schemeClr accent1 and fontRef at lt1. None of that is what we
    asked for. Overrides 1, 2 and 4 are obvious; 3 and 5 are not."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"])
    )
    shp.name = el["id"]  # so the round-trip can find it by name rather than by position
    shp.rotation = el.get("rotation", 0)  # degrees clockwise, about the shape's CENTRE

    # 1. fill -- kills fillRef -> accent1
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(el["fill"])

    # 2. line -- kills lnRef -> accent1
    stroke = el.get("stroke")
    if stroke is None:
        shp.line.fill.background()  # explicit "no line", NOT inherited
    else:
        shp.line.color.rgb = RGBColor.from_string(stroke["color"])
        # width 0 is a HAIRLINE in PowerPoint, not "no line". That is the case being measured:
        # a browser drawing border-width:0 draws nothing, and the two must be compared honestly.
        shp.line.width = Pt(stroke["width"])

    # 3. THE ONE EVERYONE FORGETS -- kills effectRef -> the Office drop shadow. No amount of
    #    fill/line overriding removes it, and there is no other shadow API.
    shp.shadow.inherit = False

    # 4. text -- kills fontRef -> lt1 (white) + theme font
    style_text_frame(shp.text_frame, el.get("text"), el.get("size_pt", 14))

    # 5. NOT in the plan's list of four, and it has to be. Overrides 1-3 win at RENDER time --
    #    an explicit <a:solidFill>/<a:ln>/<a:effectLst> in spPr beats the style ref -- but they
    #    do not DELETE the <p:style> block, so `schemeClr val="accent1"` is still sitting in the
    #    shape's XML afterwards. Plan Sec 7.3 asserts accent1 is ABSENT from every renderer-created
    #    <p:sp>; with <p:style> left in place that assertion can never pass, on a deck that is in
    #    fact rendering correctly. Strip it, and the invariant becomes checkable instead of
    #    aspirational. (<p:style> is optional in CT_Shape, so removing it is schema-valid.)
    style_el = shp._element.find(qn("p:style"))
    if style_el is not None:
        shp._element.remove(style_el)
    return shp


# --- the fixture ----------------------------------------------------------------------------
# Note the corner rectangles deliberately sit inside HEADER_RESERVED (y 0..92) and
# FOOTER_RESERVED (y 500..540). The renderer will REJECT elements there (plan Sec 3.2) and that
# is correct -- but the whole purpose of this fixture is to measure the mapping at the true slide
# corners, so the calibration deck is explicitly exempt from that rule. It is not a rendered
# document; it is a ruler.
CORNER_W, CORNER_H = 80, 40

ELEMENTS: list[dict[str, Any]] = [
    # id, x, y, w, h -- the four exact slide corners
    {"id": "corner_tl", "x": 0, "y": 0, "w": CORNER_W, "h": CORNER_H, "fill": BRAND_RED},
    {
        "id": "corner_tr",
        "x": CANVAS_W - CORNER_W,
        "y": 0,
        "w": CORNER_W,
        "h": CORNER_H,
        "fill": BRAND_RED,
    },
    {
        "id": "corner_bl",
        "x": 0,
        "y": CANVAS_H - CORNER_H,
        "w": CORNER_W,
        "h": CORNER_H,
        "fill": BRAND_RED,
    },
    {
        "id": "corner_br",
        "x": CANVAS_W - CORNER_W,
        "y": CANVAS_H - CORNER_H,
        "w": CORNER_W,
        "h": CORNER_H,
        "fill": BRAND_RED,
    },
    # the stroke ladder -- 0, 1, 4, 8 logical units
    *[
        {
            "id": f"stroke_{w}",
            "x": 60 + i * 200,
            "y": 120,
            "w": 160,
            "h": 90,
            "fill": "FFFFFF",
            "stroke": {"color": INK, "width": w},
            "text": f"stroke {w}",
        }
        for i, w in enumerate((0, 1, 4, 8))
    ],
    # rotation, about the centre
    {
        "id": "rot_30",
        "x": 60,
        "y": 250,
        "w": 120,
        "h": 80,
        "fill": BRAND_RED,
        "rotation": 30,
        "text": "30",
    },
    {
        "id": "rot_45",
        "x": 260,
        "y": 250,
        "w": 120,
        "h": 80,
        "fill": BRAND_RED,
        "rotation": 45,
        "text": "45",
    },
    # a known string in a known box: 300 units wide, minus 7.2 left + 7.2 right margin = 285.6
    # units of text width, at 14 pt. Compare the RENDERED line count against the browser's. We do
    # not predict the wrap point analytically -- Century Gothic's metrics are not ours to compute
    # -- we fix the inputs so the two renderers can be compared.
    {
        "id": "wrap_14pt",
        "x": 460,
        "y": 250,
        "w": 300,
        "h": 110,
        "fill": "FFFFFF",
        "stroke": {"color": INK, "width": 1},
        "size_pt": 14,
        "text": (
            "The quick brown fox jumps over the lazy dog and keeps running "
            "until this line has to wrap somewhere."
        ),
    },
    # mixed KO/EN -- the latin/ea split, or the lack of one
    {
        "id": "mixed_ko_en",
        "x": 60,
        "y": 380,
        "w": 400,
        "h": 80,
        "fill": "FFFFFF",
        "stroke": {"color": INK, "width": 1},
        "size_pt": 14,
        "text": "Vieworks 뷰웍스 3D Logic 개발 테스트벤치 Rev 2",
    },
]

TITLE = "Calibration — 보정 슬라이드"


def build(dst: Path) -> dict[str, Any]:
    pin = load_pin()
    prs = open_renderer_template(pin)  # fails closed if the template has drifted
    layout = prs.slide_layouts[pin["layout_signature"]["layout_index"]]
    slide = prs.slides.add_slide(layout)  # divider + logo + page number, inherited, free

    # The title goes into placeholder idx 10. On a SLIDE, placeholders[idx] is keyed by idx (it
    # raises KeyError, not IndexError) -- unlike a LAYOUT's, which is positional. See pin.py.
    try:
        ph = slide.placeholders[TITLE_PH_IDX]
    except KeyError:
        raise TemplatePinError(
            f"layout 3 did not clone placeholder idx {TITLE_PH_IDX}"
        ) from None
    tf = ph.text_frame
    tf.text = TITLE
    run = tf.paragraphs[0].runs[0]
    set_run_fonts(run)  # do NOT trust the inherited placeholder style
    run.font.size = Pt(30)
    run.font.bold = False

    expected: dict[str, dict[str, int]] = {}
    for el in ELEMENTS:
        draw_rect(slide, el)
        expected[el["id"]] = {
            "left": int(emu(el["x"])),
            "top": int(emu(el["y"])),
            "width": int(emu(el["w"])),
            "height": int(emu(el["h"])),
            "rotation": el.get("rotation", 0),
        }

    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return expected


def roundtrip(dst: Path, expected: dict[str, dict[str, int]]) -> list[str]:
    """Re-open the export and assert 0 EMU drift (plan Sec 7.3). This is the machine-checkable
    half of calibration; the raster diff is the other half and it needs a human or a renderer.

    Tolerance is 0, not "small": emu = round(u * 12700) is exact, so any drift at all means the
    write path changed the number, and a write path that rounds is a write path that will keep
    rounding.
    """
    out = Presentation(str(dst))
    problems: list[str] = []

    if len(out.slides) != 1:
        problems.append(f"expected exactly 1 slide, got {len(out.slides)}")
        return problems
    slide = out.slides[0]

    by_name = {sh.name: sh for sh in slide.shapes}
    for name, exp in expected.items():
        shp = by_name.get(name)
        if shp is None:
            problems.append(f"{name}: shape missing from the export")
            continue
        got = {
            "left": int(shp.left),
            "top": int(shp.top),
            "width": int(shp.width),
            "height": int(shp.height),
            "rotation": round(float(shp.rotation)),
        }
        for key, want in exp.items():
            if got[key] != want:
                problems.append(f"{name}.{key}: expected {want}, got {got[key]}")

        # The invariant, scoped to the shapes the RENDERER created -- not the whole slide. The
        # inherited placeholder is chrome; it is allowed to reference whatever the template says.
        xml = shp._element.xml
        if "accent1" in xml or FORBIDDEN in xml.upper():
            problems.append(
                f"{name}: accent1/{FORBIDDEN} present in the shape XML -- an add_shape override "
                "was missed (fill, line, shadow, text, or the <p:style> strip)"
            )
    return problems


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Emit the calibration fixture deck.")
    ap.add_argument("-o", "--out", type=Path, default=DST, help=f"default: {DST.name}")
    args = ap.parse_args(argv)

    try:
        expected = build(args.out)
    except TemplatePinError as exc:
        print(f"TemplatePinError: {exc}", file=sys.stderr)
        return 1

    problems = roundtrip(args.out, expected)

    sidecar = args.out.with_suffix(".json")
    sidecar.write_text(
        json.dumps(
            {
                "canvas": {"w": CANVAS_W, "h": CANVAS_H, "unit": "pt"},
                "emu_per_unit": EMU_PER_UNIT,
                "title": TITLE,
                "elements": ELEMENTS,
                "expected_emu": expected,
                "roundtrip_problems": problems,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    rel = args.out.relative_to(PROJECT_ROOT).as_posix()
    if problems:
        print(f"round-trip FAILED ({len(problems)}):", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        return 1

    print(f"wrote {rel} and {sidecar.name}: round-trip clean, 0 EMU drift on {len(expected)} shapes")
    print("Next: render it through the ACTUAL consumer and diff vs the HTML stage at zoom 1.")
    print("      Coordinate equality proves the transform, not the render (critic rule 12).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
