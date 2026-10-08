#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Logical units -> EMU, and the reserved-region predicates. Phase 1A, task 1.2.

1 logical unit == 1 pt == exactly 12,700 EMU (12192000/960 == 6858000/540 == 914400/72). The
conversion is integer and exact BY CONSTRUCTION, which is why the round-trip tolerance in
tests/canvas/test_roundtrip.py is 0 EMU and not "close enough": any drift at all means the write
path changed the number, and a write path that rounds is a write path that will keep rounding.

Pure functions. No I/O, no python-pptx beyond `Emu`. Everything downstream trusts this module, so
it stays trivially testable -- which is the whole reason it is a module and not four helpers
scattered through the renderer.

KNOWN, ACCEPTED CAVEAT -- STROKE BLEED IS NOT IN THE BBOX. `x/y/w/h` is the *path box*. A stroke is
centred on the path and bleeds `width/2` OUTWARD (SVG and DrawingML agree here, which is exactly
why the editor draws SVG). So an element sitting flush against HEADER_RESERVED with an 8-unit
stroke bleeds 4 units into it and STILL PASSES validation. Do not "fix" this by inflating the box:
the path box is what round-trips at 0 EMU, and that identity is the more valuable invariant. The
2 px raster diff (plan Sec 7.5) is what is supposed to catch real bleed.
"""

from __future__ import annotations

import math
from typing import Any

from pptx.util import Emu

EMU_PER_UNIT = 12700
CANVAS_W = 960
CANVAS_H = 540

# What a canvas-ir/1 payload must echo. A mismatch is a stale payload -- never a rescale.
CANVAS = {"w": CANVAS_W, "h": CANVAS_H, "unit": "pt"}

# Rotated bounding boxes are computed in floating point, so a corner that lands exactly on the
# slide edge can come back as 960.0000000000001. Only the ROTATED path is affected; when
# rotation == 0 the box is returned untouched and the comparison is exact.
EPS = 1e-9

# (x1, y1, x2, y2) -- top-left and bottom-right, logical units.
Box = tuple[float, float, float, float]

# --- the reserved regions (plan Sec 3.2), logical units --------------------------------------
# These exist because the logo, the gradient divider and the page number are LAYOUT-level shapes.
# After add_slide() they are not in slide.shapes at all -- they are inherited, they always render
# BEHIND slide content, and nothing you draw can be pushed under them. A collision check that
# inspects only slide.shapes will not see them, so the regions are the check.
HEADER_RESERVED: Box = (0.0, 0.0, 960.0, 92.0)  # title ph, divider picture box, logo
FOOTER_RESERVED: Box = (733.6, 500.0, 960.0, 540.0)  # slide-number chrome
BODY_SAFE: Box = (45.6, 92.0, 914.4, 500.0)  # the drawable region
TITLE_LOCKED: Box = (45.55, 26.75, 770.84, 64.0)  # not an element; its text is the `title` field

# Only these two are REJECTION regions. BODY_SAFE is the editor's snap target, not a constraint --
# an element may legally sit outside it (a full-bleed band at y 92..500, say) as long as it clears
# the header and footer. TITLE_LOCKED is not an element at all, and in any case lies wholly inside
# HEADER_RESERVED, so checking it would only ever produce a duplicate complaint.
REJECT_REGIONS: dict[str, Box] = {
    "HEADER_RESERVED": HEADER_RESERVED,
    "FOOTER_RESERVED": FOOTER_RESERVED,
}


def emu(u: float) -> Emu:
    """Logical units -> EMU. Integer and exact; the tolerance downstream is 0, not "close"."""
    return Emu(round(u * EMU_PER_UNIT))


def intersects(a: Box, b: Box) -> bool:
    """Strict overlap. HALF-OPEN: edge contact is NOT intersection.

    An element whose top edge is exactly y == 92 does not intersect HEADER_RESERVED (which ends at
    92), and it must not -- BODY_SAFE starts at exactly 92, so a top-aligned element has to be
    legal. Get this wrong and every top-aligned element in every deck is spuriously rejected.
    Likewise an element whose bottom edge is exactly y == 500 does not intersect FOOTER_RESERVED.
    """
    return a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]


def element_bbox(el: dict[str, Any]) -> Box:
    """The axis-aligned box an element actually occupies, in logical units.

    Connectors: the box spanned by their two absolute endpoints.

    Shapes and textboxes: the ROTATED AABB of the path box. The plan does not say whether a rotated
    shape is checked by its unrotated path box or by its true footprint; we check the footprint,
    because a 45-degree box that visually collides with the Vieworks logo is a real defect and the
    user-visible outcome is what the reserved regions exist to protect. When rotation == 0 this
    returns exactly the path box, so the common case is unchanged and free.

    Stroke bleed is not included -- see the module docstring.
    """
    if el.get("type") == "connector":
        x1, y1 = float(el["x1"]), float(el["y1"])
        x2, y2 = float(el["x2"]), float(el["y2"])
        return (min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2))

    x, y = float(el["x"]), float(el["y"])
    w, h = float(el["w"]), float(el["h"])
    rotation = float(el.get("rotation") or 0.0)

    if rotation % 360.0 == 0.0:
        return (x, y, x + w, y + h)

    cx, cy = x + w / 2.0, y + h / 2.0
    theta = math.radians(rotation)
    cos_t, sin_t = math.cos(theta), math.sin(theta)

    xs: list[float] = []
    ys: list[float] = []
    for px, py in ((x, y), (x + w, y), (x + w, y + h), (x, y + h)):
        dx, dy = px - cx, py - cy
        # y grows DOWNWARD, so this matrix turns CLOCKWISE on screen -- which is what DrawingML's
        # a:xfrm/@rot means. (The AABB of a centred box is symmetric in +/-theta, so a sign slip
        # here could not change the answer; the sign is stated for the reader, not for the maths.)
        xs.append(cx + dx * cos_t - dy * sin_t)
        ys.append(cy + dx * sin_t + dy * cos_t)
    return (min(xs), min(ys), max(xs), max(ys))


def reserved_violations(el: dict[str, Any]) -> list[str]:
    """Which reserved regions this element hits. Empty list == clean."""
    box = element_bbox(el)
    return [name for name, region in REJECT_REGIONS.items() if intersects(box, region)]


def out_of_bounds(el: dict[str, Any]) -> bool:
    """Does any part of the element fall outside 0,0 -> 960,540?"""
    x1, y1, x2, y2 = element_bbox(el)
    return x1 < -EPS or y1 < -EPS or x2 > CANVAS_W + EPS or y2 > CANVAS_H + EPS
