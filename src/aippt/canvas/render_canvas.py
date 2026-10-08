#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""canvas-ir/1 -> .pptx, and the CLI. Phase 1A, tasks 1.5 and 1.6.

    python -m aippt.canvas.render_canvas deck.json -o out.pptx

THE CLI IS THE POINT. The backend is fully testable and fully done before a line of editor JS is
written, and this is what you actually debug the renderer with (plan Sec 9). It exits non-zero with
the complete violation list on a bad deck.

WE ADOPT LAYOUT 3, WE DO NOT CLONE A SLIDE. The lxml slide-cloning engine in renderer_draft_brief.md
is explicitly out of scope (plan Sec 6.2). Cloning exists to REUSE HAND-DESIGNED SLIDES; the Q3
objection to add_slide() was an objection to that, and it does not bind here -- we draw our own body
content and inherit only chrome. add_slide() gives us the Vieworks logo, the gradient divider and
the page number for free, at layout level, correctly positioned, and always behind our content.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import PP_ALIGN

if __package__ in (None, ""):  # `python src/aippt/canvas/render_canvas.py`, no install
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from aippt.canvas.connectors import draw_connector  # noqa: E402
from aippt.canvas.pin import (  # noqa: E402
    TITLE_PH_IDX,
    TemplatePinError,
    load_pin,
    open_renderer_template,
)
from aippt.canvas.style import INK, add_runs, draw_shape  # noqa: E402
from aippt.canvas.validate import (  # noqa: E402
    CanvasValidationError,
    StalePayloadError,
    validate,
)

TITLE_SIZE_PT = 30

# Every run property the product depends on, stated explicitly. The layout says 30 pt and black,
# but DO NOT TRUST THE INHERITED PLACEHOLDER STYLE to match the editor's default -- say it anyway.
TITLE_RUN: dict[str, Any] = {
    "size": TITLE_SIZE_PT,
    "bold": False,
    "italic": False,
    "color": f"#{INK}",
}


def set_title(slide: Any, text: str) -> None:
    """The title -> placeholder idx 10, BY IDX, never by position.

    Layout 3 has no semantic TITLE placeholder, so slide.shapes.title is None.

    THE TWO COLLECTIONS KEY DIFFERENTLY. On a LAYOUT, placeholders[10] is POSITIONAL (layout 3 has
    exactly one placeholder, so it would raise IndexError -- pin.py uses .get(idx) for that reason).
    On a SLIDE it is keyed by IDX and raises KeyError. This runs on a slide.
    """
    try:
        ph = slide.placeholders[TITLE_PH_IDX]
    except KeyError:
        raise TemplatePinError(
            f"layout 3 did not clone placeholder idx {TITLE_PH_IDX}"
        ) from None

    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    # add_runs, not `tf.text = text`: TextFrame.text adds no run at all for an empty string, and
    # `p.runs[0]` would then raise IndexError on a deck whose title is "". It also gives the title
    # the same \n -> <a:br/> handling as body text, for free and by the same code path.
    add_runs(p, {"t": text, **TITLE_RUN})

    # The frame itself is inherited chrome -- the layout sized and placed it on-brand, and
    # TITLE_LOCKED says it is neither movable nor resizable. We correct the RUN, not the box.


def draw_image(slide: Any, el: dict[str, Any], media_map: dict[str, str] | None = None) -> Any:
    left, top = emu(el["x"]), emu(el["y"])
    width, height = emu(el["w"]), emu(el["h"])
    path = media_map.get(el["src"]) if media_map else None
    if path is None:
        from aippt.canvas.style import draw_image as legacy_draw_image
        return legacy_draw_image(slide, el)
    return slide.shapes.add_picture(path, left, top, width, height)


def draw_video(slide: Any, el: dict[str, Any], media_map: dict[str, str] | None = None) -> Any:
    left, top = emu(el["x"]), emu(el["y"])
    width, height = emu(el["w"]), emu(el["h"])
    path = media_map.get(el["src"]) if media_map else None
    if path is None:
        raise ValueError(f"video element {el['id']!r} references src {el['src']!r}, but no matching file was uploaded")
    return slide.shapes.add_movie(path, left, top, width, height, poster_frame_image=None)


def draw(slide: Any, el: dict[str, Any], media_map: dict[str, str] | None = None) -> Any:
    if el["type"] == "connector":
        return draw_connector(slide, el)
    elif el["type"] == "image":
        return draw_image(slide, el, media_map)
    elif el["type"] == "video":
        return draw_video(slide, el, media_map)
    return draw_shape(slide, el)


def render(ir: dict[str, Any], media_map: dict[str, str] | None = None) -> Presentation:
    """canvas-ir/1 -> a one-slide Presentation. Raises CanvasValidationError or TemplatePinError."""
    validate(ir)  # every violation, up front

    pin = load_pin()
    prs = open_renderer_template(pin)  # verify_pin runs inside; fails closed

    # Checked HERE and not in validate.py, so that validate.py stays pure (no pin, no filesystem)
    # and can be unit-tested without a template on disk.
    if ir["target"]["template_sha256"] != pin["renderer_template"]["sha256"]:
        raise StalePayloadError(
            f"stale payload: target.template_sha256 is "
            f"{ir['target']['template_sha256']}\n"
            f"  the pinned renderer template is "
            f"{pin['renderer_template']['sha256']}\n"
            f"The deck was authored against a different template. It is never rescaled or "
            f"re-fitted -- reload the editor and re-export."
        )

    layout = prs.slide_layouts[pin["layout_signature"]["layout_index"]]
    slide = prs.slides.add_slide(layout)  # logo + divider + page number, inherited, free

    set_title(slide, ir["title"])
    for el in ir["elements"]:  # array order IS z-order
        draw(slide, el, media_map)
    return prs


def render_append(ir: dict[str, Any], target_bytes: bytes, media_map: dict[str, str] | None = None) -> Presentation:
    """canvas-ir/1 -> appends a slide to an existing .pptx in memory."""
    validate(ir)

    pin = load_pin()

    if ir["target"]["template_sha256"] != pin["renderer_template"]["sha256"]:
        raise StalePayloadError(
            f"stale payload: target.template_sha256 is "
            f"{ir['target']['template_sha256']}\n"
            f"  the pinned renderer template is "
            f"{pin['renderer_template']['sha256']}\n"
            f"The deck was authored against a different template. It is never rescaled or "
            f"re-fitted -- reload the editor and re-export."
        )

    import io
    prs = Presentation(io.BytesIO(target_bytes))
    
    try:
        layout = prs.slide_layouts[pin["layout_signature"]["layout_index"]]
    except IndexError:
        raise TemplatePinError("Target presentation does not have the expected layout index.")

    slide = prs.slides.add_slide(layout)

    set_title(slide, ir["title"])
    for el in ir["elements"]:
        draw(slide, el, media_map)
        
    return prs


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="python -m aippt.canvas.render_canvas",
        description="Render a canvas-ir/1 deck onto the pinned Vieworks layout 3.",
    )
    ap.add_argument("deck", type=Path, help="canvas-ir/1 JSON")
    ap.add_argument("-o", "--out", type=Path, required=True, help="destination .pptx")
    args = ap.parse_args(argv)

    try:
        ir = json.loads(args.deck.read_text(encoding="utf-8"))
    except FileNotFoundError:
        print(f"deck not found: {args.deck}", file=sys.stderr)
        return 2
    except json.JSONDecodeError as exc:
        print(f"{args.deck} is not valid JSON: {exc}", file=sys.stderr)
        return 2

    try:
        prs = render(ir)
    except CanvasValidationError as exc:  # StalePayloadError is a subclass
        print(f"{args.deck}: {exc}", file=sys.stderr)
        return 1
    except TemplatePinError as exc:
        print(f"TemplatePinError: {exc}", file=sys.stderr)
        return 1

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(f"wrote {args.out} -- 1 slide, {len(ir['elements'])} element(s) on layout 3")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
