#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fonts, text frames and the add_shape overrides. Phase 1A, task 1.3.

add_shape() does not hand you a blank shape. `new_autoshape_sp` hard-codes a <p:style> block
pointing lnRef, fillRef and effectRef at `schemeClr accent1` and fontRef at `lt1` -- and accent1 in
THIS template is #4472C4, stock Office blue (reference/layout_manifest.json). So every shape
arrives blue-filled, blue-outlined, drop-shadowed, with white theme-font text. None of that is what
we asked for, and none of it goes away by accident.

There are FIVE overrides, not the four the plan lists. Phase 0 found the fifth
(phase0_change_detail.md:57) and tools/calibrate.py:168 has it working:

    1. fill          kills fillRef   -> accent1
    2. line          kills lnRef     -> accent1
    3. shadow        kills effectRef -> the Office drop shadow    <- the one everyone forgets
    4. text          kills fontRef   -> lt1 (white) + theme font
    5. strip <p:style>                                            <- not in the plan, and required

Overrides 1-3 win at RENDER time -- an explicit <a:solidFill>/<a:ln>/<a:effectLst> in spPr beats
the style ref -- but they do NOT delete the <p:style> block, so `schemeClr val="accent1"` is still
sitting in the shape's XML afterwards. Plan Sec 7.3 asserts accent1 is ABSENT from every
renderer-created <p:sp>; with <p:style> left in place that assertion can never pass, on a deck that
is in fact rendering correctly. Stripping it is what makes the invariant checkable rather than
aspirational. (<p:style> is optional in CT_Shape, so removing it is schema-valid.)

NOTE ON DUPLICATION: tools/calibrate.py inlines its own copy of emu(), set_run_fonts() and
style_text_frame(). That is deliberate and it stays for now -- Phase 0's gates have to keep passing
standalone. De-duping it re-opens a passing gate and is a separate, reviewed change.
"""

from __future__ import annotations

import base64
import io
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Pt

from aippt.canvas.geometry import emu

LATIN = "Century Gothic"
EA = "맑은 고딕"  # Century Gothic ships NO Hangul glyphs. This is not optional.

BRAND_RED = "FF0032"  # the default for new content
INK = "000000"
FORBIDDEN = "4472C4"  # theme accent1. If this reaches the output, an override was missed.

# D-1 IS STILL OPEN. Century Gothic is NOT the Vieworks brand font -- the template's real design
# fonts are 뷰웍스 B / 뷰웍스 M, set at run level on every exemplar slide, so Century Gothic body
# text will sit next to inherited 뷰웍스 chrome and look subtly off-brand. We ship Century Gothic
# because the user asked for it. Today that is ONE CONSTANT here (and one CSS declaration in the
# Phase 1B editor); after the editor ships it is an expensive discovery. Raise it before Phase 1
# closes. See plan Sec 10, D-1.

PRST: dict[str, MSO_SHAPE] = {
    "rect": MSO_SHAPE.RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
}

# Mapped with an EXPLICIT dict, never PP_ALIGN[name]. Enum-member-name indexing happens to work in
# python-pptx 1.0.x, but it couples the IR's public vocabulary to an upstream enum's member names,
# and it yields a bare KeyError where the caller wanted a validation error.
ALIGN: dict[str, PP_ALIGN] = {
    "LEFT": PP_ALIGN.LEFT,
    "CENTER": PP_ALIGN.CENTER,
    "RIGHT": PP_ALIGN.RIGHT,
}
VALIGN: dict[str, MSO_ANCHOR] = {
    "TOP": MSO_ANCHOR.TOP,
    "MIDDLE": MSO_ANCHOR.MIDDLE,
    "BOTTOM": MSO_ANCHOR.BOTTOM,
}

DEFAULT_ALIGN = "CENTER"
DEFAULT_VALIGN = "MIDDLE"


def rgb(color: str) -> RGBColor:
    """'#FF0032' or 'FF0032' -> RGBColor. The IR always carries the '#'; python-pptx never wants it."""
    color = color.lstrip("#")
    if len(color) == 8:
        color = color[:6]  # Drop the alpha channel for python-pptx compatibility
    return RGBColor.from_string(color)


def apply_color(color_format: Any, color: str) -> None:
    """Parses '#RRGGBB' or '#RRGGBBAA' and applies it to the color_format.
    Sets RGB value and inserts <a:alpha> element if an alpha channel is present.
    """
    color = color.lstrip("#")
    
    rgb_hex = color[:6]
    color_format.rgb = RGBColor.from_string(rgb_hex)
    
    if len(color) == 8:
        alpha_hex = color[6:8]
        alpha_pct = int(round((int(alpha_hex, 16) / 255.0) * 100000))
        xClr = color_format._color._xClr
        if xClr is not None:
            existing = xClr.find(qn("a:alpha"))
            if existing is not None:
                xClr.remove(existing)
            alpha_el = parse_xml(f'<a:alpha xmlns:a="{qn("a:alpha").split("}")[0][1:]}" val="{alpha_pct}"/>')
            xClr.append(alpha_el)


def set_run_fonts(run: Any, latin: str = LATIN, ea: str = EA) -> None:
    """The latin/EA split. Copied verbatim from tools/calibrate.py, where it has actually run.

    python-pptx cannot express an East-Asian typeface: Font.name writes only <a:latin>, and
    CT_TextCharacterProperties declares no `ea` child at all. So <a:ea> has to be inserted by hand.

    The successor list is the DrawingML child order for CT_TextCharacterProperties: <a:ea> must
    follow <a:latin> and precede <a:cs>/<a:sym>/<a:hlink*>/<a:rtl>/<a:extLst>. IT MUST STAY
    COMPLETE, not merely plausible -- insert_element_before APPENDS when it finds none of its
    successors, so a missing entry silently puts <a:ea> in the wrong place and PowerPoint calls the
    file corrupt.
    """
    rPr = run._r.get_or_add_rPr()
    rPr.get_or_add_latin().typeface = latin
    ea_el = rPr.find(qn("a:ea"))
    if ea_el is None:
        ea_el = rPr.makeelement(qn("a:ea"), {})
        rPr.insert_element_before(
            ea_el, "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"
        )
    ea_el.set("typeface", ea)


def add_runs(p: Any, r: dict[str, Any]) -> None:
    """One IR run -> one or more DrawingML runs, separated by <a:br/>.

    Plan Sec 5.6 maps the editor's <br> to <a:br/> ("line break within a paragraph"), but the Sec 11
    appendix gives no way to express one. Decision: permit '\\n' inside a run's `t` and split here,
    carrying the run's formatting across the break. This keeps the run model flat and needs no
    schema change beyond documenting it.

    The split is not cosmetic. A raw '\\n' left inside <a:t> is only whitespace to DrawingML -- it
    renders as a space, not a break -- so splitting is the only thing that makes the character mean
    what the author typed. (python-pptx would not escape it either: its control-character escape
    covers \\x00-\\x08 and \\x0B-\\x1F, and deliberately not \\t or \\n.)

    U+00A0 (nbsp) is passed through untouched and is never collapsed to a space.
    """
    for i, piece in enumerate(r["t"].split("\n")):
        if i:
            p.add_line_break()  # -> <a:br/>, in document order between the two runs
        run = p.add_run()
        run.text = piece
        set_run_fonts(run)
        run.font.size = Pt(r["size"])  # 1:1 with the model. No factor.
        run.font.bold = bool(r.get("bold", False))
        run.font.italic = bool(r.get("italic", False))
        apply_color(run.font.color, r["color"])


def style_text_frame(tf: Any, text: dict[str, Any] | None) -> None:
    """Override 4, and THE MARGIN TRAP.

    PowerPoint defaults a text frame to 0.1 in left/right and 0.05 in top/bottom -- 7.2 and 3.6
    logical units. If the browser's div uses padding:0 while PowerPoint keeps its defaults, the two
    wrap text at a DIFFERENT WIDTH, and the result is an off-by-one-line overflow that presents as
    a coordinate bug. Set all four explicitly. These mirror the CSS contract in plan Sec 5.6 exactly
    and must only ever be changed together with it.

    auto_size = NONE, always. TextFrame.fit_text() does exist, but it needs a TTF path, it is not a
    PowerPoint rendering engine, and it is a poor foundation for mixed-font CJK. Do not build on it
    (plan Sec 6.6). The human sizes the box; rendered QA is the authority.
    """
    tf.margin_left = tf.margin_right = Pt(7.2)
    tf.margin_top = tf.margin_bottom = Pt(3.6)
    tf.word_wrap = True  # == CSS white-space: pre-wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = VALIGN[(text or {}).get("valign", DEFAULT_VALIGN)]

    if not text:
        return

    alignment = ALIGN[text.get("align", DEFAULT_ALIGN)]
    tf.clear()  # leaves exactly one empty paragraph behind
    for i, para in enumerate(text["paragraphs"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alignment
        p.space_before = Pt(0)  # explicit -- do not inherit
        p.space_after = Pt(0)
        p.line_spacing = 1.2  # == CSS line-height: 1.2
        for r in para["runs"]:
            add_runs(p, r)


def strip_style(shp: Any) -> None:
    """Override 5. Applies to <p:sp> AND <p:cxnSp> -- a connector carries a <p:style> block for
    exactly the same reason a shape does. See the module docstring for why this is not optional."""
    style_el = shp._element.find(qn("p:style"))
    if style_el is not None:
        shp._element.remove(style_el)


def apply_fill(shp: Any, fill: str | None) -> None:
    """Override 1. `None` means NO fill -- an explicit <a:noFill/>, which still beats fillRef."""
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        apply_color(shp.fill.fore_color, fill)


def apply_line(shp: Any, stroke: dict[str, Any] | None) -> None:
    """Override 2.

    `None` means no line, and that has to be said explicitly -- never inherited. Width 0 is a
    HAIRLINE in PowerPoint, not "no line": a browser draws nothing for border-width:0 while
    PowerPoint draws the thinnest line it can. So a 0-width stroke is a legal, meaningful value and
    it is NOT the way to ask for no line.
    """
    if stroke is None:
        shp.line.fill.background()
    else:
        apply_color(shp.line.color, stroke["color"])
        shp.line.width = Pt(stroke["width"])


def draw_shape(slide: Any, el: dict[str, Any]) -> Any:
    """rect | ellipse | triangle | textbox -> a styled shape on `slide`. All five overrides."""
    etype = el["type"]
    left, top = emu(el["x"]), emu(el["y"])
    width, height = emu(el["w"]), emu(el["h"])

    if etype == "textbox":
        # add_textbox() carries no <p:style> and so inherits no accent1 -- but it DOES ship
        # wrap="none" and <a:spAutoFit/>, both of which style_text_frame overrides below. The
        # remaining overrides are applied anyway, so that every renderer-created shape satisfies
        # the same assertions regardless of how it was born.
        shp = slide.shapes.add_textbox(left, top, width, height)
    else:
        shp = slide.shapes.add_shape(PRST[etype], left, top, width, height)

    shp.name = el["id"]  # the round-trip suite finds shapes by NAME, not by position
    shp.rotation = float(el.get("rotation") or 0.0)  # degrees CW, about the shape's CENTRE

    apply_fill(shp, el.get("fill"))
    apply_line(shp, el.get("stroke"))
    shp.shadow.inherit = False  # 3 -- no amount of fill/line overriding removes the drop shadow
    style_text_frame(shp.text_frame, el.get("text"))
    strip_style(shp)
    return shp


def draw_image(slide: Any, el: dict[str, Any]) -> Any:
    """image -> a picture on `slide`. `src` is a Base64 data: URI, decoded straight into memory --
    the editor sends the bytes inline; there is no upload directory or filename to trust."""
    left, top = emu(el["x"]), emu(el["y"])
    width, height = emu(el["w"]), emu(el["h"])

    src = el.get("src", "")
    if src.startswith("data:"):
        # format: data:image/png;base64,iVBOR...
        header, b64 = src.split(",", 1)
        img_data = base64.b64decode(b64)
        img_io = io.BytesIO(img_data)
        shp = slide.shapes.add_picture(img_io, left, top, width, height)
    else:
        raise ValueError(f"Unsupported image src format for element {el['id']}")

    shp.name = el["id"]
    shp.rotation = float(el.get("rotation") or 0.0)
    return shp
