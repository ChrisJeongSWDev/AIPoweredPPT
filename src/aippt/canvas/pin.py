#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Template pinning. Two gates, both mandatory, both at PROCESS START -- not per request.

Hard-coding `slide_layouts[3]` is safe ONLY while the package is pinned (critic rule 4). A
template that is changed, re-cut or re-ordered must fail closed with a clear error. It must never
silently select a different layout, because a "successful" export onto the wrong layout is a
worse outcome than a crash: it looks fine, ships, and is discovered by the audience.

    pin = load_pin()
    verify_pin(pin)                      # raises TemplatePinError, or returns
    prs = open_renderer_template(pin)    # verified again; cheap next to being wrong

Gate 1 is byte identity of the renderer template. Gate 2 is its structural signature, and it
exists because gate 1 alone can be defeated by the one mistake we actually expect: someone
rebuilds the template, sees the SHA mismatch, pastes the new SHA into template_pin.json to make
the error go away, and never notices that layout 3 is now a different layout.

There is no auto-repin, no --force, and no environment-variable escape hatch. Re-pinning means
running the builder, reading its output, and committing the new artifact deliberately (Sec 4.3).

This module never reads the SOURCE template: it is a build-time input, not a runtime one, and the
request path must not depend on a 1.8 MB deck it does not use.

This module exposes exactly three functions -- load_pin, verify_pin, open_renderer_template. The
hashing, path resolution and signature checks are deliberately INLINED into them rather than
split into helpers: the surface is the contract, and a helper is one more thing a caller can
reach for instead of the gates.
"""

from __future__ import annotations

import hashlib
import json
import posixpath
from pathlib import Path
from typing import Any

from pptx import Presentation

# <root>/src/aippt/canvas/pin.py -> <root>. The pin file names paths relative to the project
# root, so they must be resolved against it and not against the caller's cwd.
PROJECT_ROOT = Path(__file__).resolve().parents[3]
PIN_PATH = PROJECT_ROOT / "reference" / "template_pin.json"

PIN_VERSION = 1
TITLE_PH_IDX = 10

# What make_renderer_template.py has not yet filled in. A freshly-checked-out pin carries this
# and MUST fail: an unpinned template is an unverified template.
UNPINNED = "<FILL FROM make_renderer_template.py OUTPUT>"

_REBUILD = (
    "Rebuild via `python src/aippt/tools/make_renderer_template.py --write-pin` "
    "and RE-PIN deliberately."
)


class TemplatePinError(RuntimeError):
    """Fail closed. Never fall back to another layout."""


def load_pin() -> dict[str, Any]:
    """Read and structurally validate template_pin.json itself.

    The path is PIN_PATH, not a parameter: there is exactly one pin per checkout, and letting a
    caller point this at some other file is the escape hatch the module docstring says we do not
    have.

    This only checks that the pin is a well-formed pin. It says nothing about whether the package
    on disk matches it -- that is verify_pin's job, and keeping them apart is what lets a caller
    report "your pin file is malformed" separately from "your template has drifted".
    """
    try:
        pin = json.loads(PIN_PATH.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise TemplatePinError(f"pin file not found: {PIN_PATH}\n{_REBUILD}") from None
    except json.JSONDecodeError as exc:
        raise TemplatePinError(f"pin file is not valid JSON: {PIN_PATH}\n  {exc}") from None

    if pin.get("pin_version") != PIN_VERSION:
        raise TemplatePinError(
            f"pin_version is {pin.get('pin_version')!r}, this code speaks {PIN_VERSION}"
        )

    for key in ("source_template", "renderer_template", "layout_signature"):
        if key not in pin:
            raise TemplatePinError(f"pin file is missing required key {key!r}")

    rt, sig = pin["renderer_template"], pin["layout_signature"]
    for key in ("filename", "sha256", "bytes", "built_from_sha256"):
        if key not in rt:
            raise TemplatePinError(f"pin file is missing renderer_template.{key!r}")
    for key in (
        "layout_index",
        "layout_partname",
        "slide_width_emu",
        "slide_height_emu",
        "slide_count",
        "layout_count",
        "placeholder_idxs",
        f"placeholder_{TITLE_PH_IDX}_emu",
        "required_media",
    ):
        if key not in sig:
            raise TemplatePinError(f"pin file is missing layout_signature.{key!r}")

    # The pin must be internally consistent before we trust it about anything else.
    if rt["built_from_sha256"] != pin["source_template"].get("sha256"):
        raise TemplatePinError(
            "pin file is self-inconsistent: renderer_template.built_from_sha256 does not match "
            "source_template.sha256. Someone edited it by hand.\n" + _REBUILD
        )
    return pin


def verify_pin(pin: dict[str, Any]) -> None:
    """Both gates. Returns None; raises TemplatePinError on any mismatch."""
    tpl = PROJECT_ROOT / pin["renderer_template"]["filename"]

    # --- Gate 1: byte identity -------------------------------------------------------------
    expected = pin["renderer_template"]["sha256"]
    if expected == UNPINNED:
        raise TemplatePinError(
            f"the renderer template has never been pinned ({tpl.name} sha256 is still the "
            f"placeholder).\n{_REBUILD}"
        )
    if not tpl.is_file():
        raise TemplatePinError(f"renderer template not found: {tpl}\n{_REBUILD}")

    actual_bytes = tpl.stat().st_size
    if actual_bytes != pin["renderer_template"]["bytes"]:
        raise TemplatePinError(
            f"renderer template size mismatch\n  expected {pin['renderer_template']['bytes']} B"
            f"\n  actual   {actual_bytes} B\n{_REBUILD}"
        )

    # Streamed in 1 MiB chunks: the deck is small today, but a hash that needs the whole file
    # resident is a footgun waiting for the day it is not.
    digest = hashlib.sha256()
    with tpl.open("rb") as f:
        while chunk := f.read(1 << 20):
            digest.update(chunk)
    actual = digest.hexdigest()
    if actual != expected:
        raise TemplatePinError(
            f"renderer template SHA mismatch\n  expected {expected}\n  actual   {actual}\n"
            f"{_REBUILD}"
        )

    # --- Gate 2: structural signature -------------------------------------------------------
    # Catches the re-pinned-but-wrong package: right bytes for a pin someone pasted in, wrong deck.
    prs = Presentation(str(tpl))
    sig = pin["layout_signature"]

    checks: list[tuple[bool, str]] = [
        (len(prs.slides) == sig["slide_count"], "renderer template must contain 0 slides"),
        (len(prs.slide_layouts) == sig["layout_count"], "layout count changed"),
        (prs.slide_width == sig["slide_width_emu"], "slide width changed"),
        (prs.slide_height == sig["slide_height_emu"], "slide height changed"),
    ]

    idx = sig["layout_index"]
    if not 0 <= idx < len(prs.slide_layouts):
        raise TemplatePinError(f"layout signature mismatch: no layout at index {idx}")
    layout = prs.slide_layouts[idx]

    checks += [
        (str(layout.part.partname) == sig["layout_partname"], f"layout {idx} partname changed"),
        (
            sorted(p.placeholder_format.idx for p in layout.placeholders)
            == sig["placeholder_idxs"],
            f"layout {idx} placeholder idx set changed",
        ),
    ]

    got = {
        posixpath.basename(str(rel.target_partname))
        for rel in layout.part.rels.values()
        if not rel.is_external
    }
    missing = sorted(set(sig["required_media"]) - got)
    checks.append((not missing, f"layout {idx} lost media {missing}"))

    # .get(idx) -- on a LAYOUT, placeholders[10] is POSITIONAL (layout 3 has one placeholder, so
    # it would raise IndexError). Only slide.placeholders is keyed by idx. See plan Sec 6.3.
    ph = layout.placeholders.get(TITLE_PH_IDX)
    if ph is None:
        raise TemplatePinError(
            f"layout signature mismatch: layout {idx} has no placeholder idx {TITLE_PH_IDX}"
        )
    exp = sig[f"placeholder_{TITLE_PH_IDX}_emu"]
    checks.append(
        (
            (ph.left, ph.top, ph.width, ph.height)
            == (exp["left"], exp["top"], exp["width"], exp["height"]),
            f"placeholder idx {TITLE_PH_IDX} geometry moved: "
            f"expected {exp}, got left={ph.left} top={ph.top} width={ph.width} height={ph.height}",
        )
    )

    for ok, msg in checks:
        if not ok:
            raise TemplatePinError(f"layout signature mismatch: {msg}\n{_REBUILD}")


def open_renderer_template(pin: dict[str, Any]) -> Presentation:
    """The ONLY sanctioned way to obtain the renderer's Presentation. Verified, then opened.

    Yes, this opens the package twice (verify_pin opens one to check the signature, we open
    another to hand back). That is a few milliseconds once per process, and the alternative --
    handing back the instance verify_pin happened to build -- couples the two functions for no
    benefit. The brief asks for verify_pin() -> None; this keeps that honest.
    """
    verify_pin(pin)
    return Presentation(str(PROJECT_ROOT / pin["renderer_template"]["filename"]))


if __name__ == "__main__":
    # `python src/aippt/canvas/pin.py` -- the fail-closed check, on demand.
    import sys

    try:
        p = load_pin()
        verify_pin(p)
    except TemplatePinError as exc:
        print(f"TemplatePinError: {exc}", file=sys.stderr)
        raise SystemExit(1) from None
    print(f"pin OK: {p['renderer_template']['filename']} @ {p['renderer_template']['sha256']}")
