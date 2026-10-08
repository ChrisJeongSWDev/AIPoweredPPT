#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""slot_id -> shape_id -> the shape. And how to write into it without destroying the design.

Why shape_id and not shape.name: layout `1_사용자 지정 레이아웃` has three DIFFERENT placeholders
at three different positions with three different font sizes (28/36/16pt) and all three are
named `텍스트 개체 틀 8` (layout_manifest.json:161,195,225). PowerPoint's auto-generated Korean
names are neither semantic nor unique. cNvPr/@id is unique within a slide, stable across
renaming and re-ordering, and broken only when someone deletes and recreates a shape -- which
is exactly when you WANT a hard failure.
"""

from __future__ import annotations

import copy
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from .brand import style_code_run

_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

#: Geometry fingerprints are compared with this slack. 1000 EMU is ~0.001 inch -- far below
#: anything a human could have moved a shape by, and far above float noise.
GEOMETRY_TOLERANCE_EMU = 1000


class ExemplarManifestNotCurated(RuntimeError):
    """The manifest exists but no human has assigned slot_ids yet."""


class TemplateDriftError(RuntimeError):
    """A slot's fingerprint no longer matches the template. Never guess past this."""


@dataclass
class Slot:
    slot_id: str
    kind: str                      # text | picture | table | group
    shape_id: int                  # PRIMARY KEY
    shape_name: str                # fingerprint only -- NOT unique, never a key
    z_index: int                   # fingerprint only
    emu: dict[str, int]
    effective_style: dict[str, Any] | None = None
    budget: dict[str, Any] | None = None
    sample_text: str | None = None
    required: bool = False
    on_unbound: str = "remove"     # remove | keep | blank


@dataclass
class Exemplar:
    exemplar_id: str
    source_slide_index: int
    layout_name: str
    roles: list[str] = field(default_factory=list)
    supports_blocks: list[str] = field(default_factory=list)
    slots: dict[str, Slot] = field(default_factory=dict)
    brand_shapes: list[int] = field(default_factory=list)


def load_exemplar_manifest(path: Path, *, template_sha256: str) -> dict[str, Exemplar]:
    """Load and pin. Refuses an uncurated or stale manifest -- loudly, at load time.

    This is the gate that makes "not yet curated" safe: nothing downstream can render against
    invented slot bindings, because there is nothing to load.
    """
    raw = json.loads(Path(path).read_text(encoding="utf-8"))

    if not raw.get("curated"):
        raise ExemplarManifestNotCurated(
            f"{path} has not been curated. Run `python -m aippt.analyze_exemplars --dump`, "
            f"read the inventory, assign slot_ids and exemplar_ids, then set \"curated\": true."
        )
    if raw["source"]["sha256"] != template_sha256:
        raise TemplateDriftError(
            f"{path} was generated against template sha256 {raw['source']['sha256'][:12]}… but "
            f"the template on disk is {template_sha256[:12]}…. Re-run analyze_exemplars.py; do "
            f"not render against a manifest that does not describe this file."
        )

    out: dict[str, Exemplar] = {}
    for ex in raw["exemplars"]:
        if not ex.get("exemplar_id"):
            continue
        slots = {s["slot_id"]: Slot(**{k: s[k] for k in
                                       ("slot_id", "kind", "shape_id", "shape_name",
                                        "z_index", "emu", "effective_style", "budget",
                                        "sample_text", "required", "on_unbound")})
                 for s in ex["slots"] if s.get("slot_id")}
        out[ex["exemplar_id"]] = Exemplar(
            exemplar_id=ex["exemplar_id"], source_slide_index=ex["source_slide_index"],
            layout_name=ex["layout_name"], roles=ex.get("roles", []),
            supports_blocks=ex.get("supports_blocks", []), slots=slots,
            brand_shapes=ex.get("brand_shapes", []))
    return out


def resolve_shape(slide: Any, slot: Slot) -> Any:
    """slot -> the shape on the cloned slide. Verifies the fingerprint. Never falls back.

    A silent substitution here yields a deck that looks fine and is wrong -- the worst possible
    outcome, because nobody notices.
    """
    for z, shape in enumerate(slide.shapes):
        if shape.shape_id != slot.shape_id:
            continue
        drift = []
        if shape.name != slot.shape_name:
            drift.append(f"name {shape.name!r} != {slot.shape_name!r}")
        if z != slot.z_index:
            drift.append(f"z_index {z} != {slot.z_index}")
        for key in ("left", "top", "width", "height"):
            actual, expected = getattr(shape, key), slot.emu.get(key)
            if actual is None or expected is None:
                continue
            if abs(int(actual) - expected) > GEOMETRY_TOLERANCE_EMU:
                drift.append(f"{key} {int(actual)} != {expected} "
                             f"(Δ{int(actual) - expected} EMU)")
        if drift:
            raise TemplateDriftError(
                f"slot {slot.slot_id!r} (shape_id {slot.shape_id}): the template has changed "
                f"under the manifest -- {'; '.join(drift)}. Re-run analyze_exemplars.py and "
                f"re-check the curation."
            )
        return shape
    raise TemplateDriftError(
        f"slot {slot.slot_id!r}: no shape with cNvPr id {slot.shape_id} on the cloned slide. "
        f"It was deleted and recreated in the template. Re-run analyze_exemplars.py."
    )


# --------------------------------------------------------------------------------------
# C-3 -- writing text without destroying the design
# --------------------------------------------------------------------------------------
#
# tf.text = "x"  -> drops every paragraph and every <a:rPr>. The 30pt 뷰웍스 B in brand red
#                   becomes 18pt Calibri black. BANNED. Grep for it in review.
# run.text = "x" -> replaces the literal text inside an existing <a:r>, leaving its <a:rPr>
#                   untouched. THIS is the primitive.

def set_text_preserving_style(tf: Any, text: str) -> None:
    """Single-line slot (title, card title, step title). Keeps run[0]'s rPr -- the design."""
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        para.add_run().text = text          # no run to inherit from: last resort
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def rewrite_paragraphs(tf: Any, lines: list[tuple[int, Any]]) -> None:
    """Multi-paragraph slot (bullets, body prose). ``lines`` is [(level, RichText), ...].

    Harvest a style PROTOTYPE from the exemplar before touching anything, then stamp it per
    line. The prototype carries <a:pPr> (indent, bullet char, spacing) and <a:rPr> (size,
    typeface, colour) -- i.e. it carries the design.
    """
    txBody = tf._txBody
    protos: dict[int, Any] = {}
    for para in tf.paragraphs:
        protos.setdefault(para.level, copy.deepcopy(para._p))
    proto0 = protos.get(0) or copy.deepcopy(tf.paragraphs[0]._p)

    for para in list(tf.paragraphs):
        txBody.remove(para._p)

    for level, rich in lines:
        new_p = copy.deepcopy(protos.get(level, proto0))
        _restructure_runs(new_p, rich)
        txBody.append(new_p)

    # Second pass, over the TYPED proxies: python-pptx inserts rPr children in schema order,
    # which is fiddly to get right by hand (a:solidFill before a:latin before a:hlinkClick...).
    for para, (level, rich) in zip(tf.paragraphs, lines):
        para.level = level                                  # explicit, in case we reused proto0
        for run, ir_run in zip(para.runs, rich.runs):
            if ir_run.b:
                run.font.bold = True
            if ir_run.i:
                run.font.italic = True
            if ir_run.code:
                style_code_run(run)                          # the one sanctioned deviation
            if ir_run.link:
                run.hyperlink.address = ir_run.link


def _restructure_runs(p_element: Any, rich: Any) -> None:
    """Keep <a:pPr> and <a:endParaRPr>; drop the existing <a:r>/<a:br>/<a:fld>; append one
    <a:r> per IR run, each carrying a deepcopy of the prototype's <a:rPr>."""
    proto_rPr = None
    first_r = p_element.find(qn("a:r"))
    if first_r is not None and first_r.find(qn("a:rPr")) is not None:
        proto_rPr = copy.deepcopy(first_r.find(qn("a:rPr")))
    elif (end := p_element.find(qn("a:endParaRPr"))) is not None:
        proto_rPr = copy.deepcopy(end)
        proto_rPr.tag = qn("a:rPr")          # endParaRPr and rPr share CT_TextCharacterProperties

    for tag in ("a:r", "a:br", "a:fld"):
        for el in p_element.findall(qn(tag)):
            p_element.remove(el)

    end_para = p_element.find(qn("a:endParaRPr"))
    for ir_run in rich.runs:
        r = parse_xml("<a:r %s><a:t/></a:r>" % nsdecls("a"))
        if proto_rPr is not None:
            r.insert(0, copy.deepcopy(proto_rPr))
        t = r.find(qn("a:t"))
        t.text = ir_run.t
        if ir_run.t != ir_run.t.strip():
            t.set(_XML_SPACE, "preserve")    # leading/trailing space is significant
        if end_para is not None:
            end_para.addprevious(r)
        else:
            p_element.append(r)


# --------------------------------------------------------------------------------------
# §4.3 -- images: swap the blip, keep the frame
# --------------------------------------------------------------------------------------

def replace_picture(slide: Any, pic: Any, image_stream: Any) -> Any:
    """Swap the image PART underneath the existing <p:pic>.

    The exemplar's picture already has the right position, size, and often a picture style
    (rounded corners, a border, a shadow). Deleting it and calling add_picture() at the same
    geometry throws all of that away.
    """
    image_part, rId = slide.part.get_or_add_image_part(image_stream)
    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.set(qn("r:embed"), rId)
    _fit_src_rect(pic, image_part)
    return pic


def _fit_src_rect(pic: Any, image_part: Any) -> None:
    """The trap: the new image is now stretched to the old frame's aspect ratio.

    PowerPoint's own behaviour is aspect-FILL with a centre crop, expressed as <a:srcRect>
    inside <a:blipFill>. Units are thousandths of a percent (50000 == 50%).

    Pillow is a HARD dependency precisely here: without image.size, aspect-fill degrades
    silently into stretch.
    """
    px_w, px_h = image_part.image.size
    if not px_w or not px_h or not pic.width or not pic.height:
        return
    frame_ar, img_ar = pic.width / pic.height, px_w / px_h

    if img_ar > frame_ar:                     # image too wide -> crop left/right
        crop = (1 - frame_ar / img_ar) / 2
        attrs = {"l": str(int(crop * 100000)), "r": str(int(crop * 100000))}
    else:                                     # image too tall -> crop top/bottom
        crop = (1 - img_ar / frame_ar) / 2
        attrs = {"t": str(int(crop * 100000)), "b": str(int(crop * 100000))}

    blip_fill = pic._element.blipFill
    for old in blip_fill.findall(qn("a:srcRect")):
        blip_fill.remove(old)                 # reset any inherited crop first
    src_rect = parse_xml("<a:srcRect %s/>" % nsdecls("a"))
    for k, v in attrs.items():
        src_rect.set(k, v)
    # Child order in CT_BlipFillProperties: a:blip, a:srcRect, then a:tile|a:stretch.
    blip_fill.find(qn("a:blip")).addnext(src_rect)


# --------------------------------------------------------------------------------------
# C-4 -- the leftover-demo-text problem
# --------------------------------------------------------------------------------------

def remove_unbound_slots(slide: Any, exemplar: Exemplar, bound: set[str]) -> int:
    """Every slot no IR block bound MUST go.

    Skip this and the exemplar's own demo content -- "제목을 입력하세요", a sample bullet, a stock
    photo -- ships inside the customer's deck, on brand, in the right font, looking entirely
    intentional. This is THE signature failure of slide cloning, and a shape-count assertion
    will not catch it.

    `decorations` are never touched. That is the entire point of cloning.
    """
    removed = 0
    for slot in exemplar.slots.values():
        if slot.slot_id in bound:
            continue
        if slot.required:
            raise ValueError(
                f"exemplar {exemplar.exemplar_id!r} requires slot {slot.slot_id!r}, but no "
                f"block bound it."
            )
        shape = resolve_shape(slide, slot)
        if slot.on_unbound == "remove":
            shape._element.getparent().remove(shape._element)
            removed += 1
        elif slot.on_unbound == "blank" and shape.has_text_frame:
            shape.text_frame.clear()      # keeps the box -- e.g. a design element that must stay
    return removed