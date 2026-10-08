from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(255, 0, 0)
xClr = shape.fill.fore_color._color._xClr

alpha_pct = 50196
alpha_el = parse_xml(f'<a:alpha xmlns:a="{qn("a:alpha").split("}")[0][1:]}" val="{alpha_pct}"/>')
xClr.append(alpha_el)

prs.save("test_alpha_save.pptx")
print("Saved successfully.")
