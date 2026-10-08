from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
import xml.etree.ElementTree as ET
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(255, 0, 0)
srgbClr = shape.fill.fore_color._color._xClr
print(type(srgbClr))
print(ET.tostring(srgbClr).decode("utf-8"))

alpha = parse_xml(f'<a:alpha xmlns:a="{qn("a:alpha").split("}")[0][1:]}" val="50000"/>')
srgbClr.append(alpha)

print(ET.tostring(srgbClr).decode("utf-8"))
