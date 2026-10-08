#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""GATE 2 -- the editor server. Phase 1B, plan Sec 9 tasks 3.1-3.3.

FastAPI's TestClient. No browser, no uvicorn, no port.

app.py is an ADAPTER, so what is under test here is the ADAPTATION, not the rendering: does a valid
payload come back as a .pptx, and does each way of being invalid come back as the status the editor
knows how to act on? The renderer's own correctness is tests/canvas/test_roundtrip.py's job and is
not re-litigated here.

THE ONE THAT MATTERS MOST is test_stale_sha_is_409 / test_stale_canvas_is_409. StalePayloadError
SUBCLASSES CanvasValidationError, so an `except CanvasValidationError` written above the stale
handler would swallow it and report every stale payload as a 422 -- telling the user to fix a deck
that is perfectly fine, when what they actually need to do is reload. The ordering in app.py is
load-bearing, it is invisible, and these two tests are what hold it in place.

REQUIRES THE EDITOR EXTRA:  pip install -e ".[editor,dev]"
Without fastapi + httpx this module SKIPS rather than fails, so that `pytest tests/` stays green for
someone who installed only the renderer -- the CLI is a first-class path and must not need a web
stack. If you are running Gate 2, check the summary says PASSED and not SKIPPED.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from io import BytesIO
from pathlib import Path

import pytest

_EXTRA = "the editor extra is not installed -- run: pip install -e \".[editor,dev]\""

pytest.importorskip("fastapi", reason=_EXTRA)
pytest.importorskip("httpx", reason=_EXTRA)  # FastAPI's TestClient is built on httpx

from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402

from aippt.canvas.app import MOCK_LATENCY_S, app, mock_ir  # noqa: E402
from aippt.canvas.pin import load_pin  # noqa: E402
from aippt.canvas.validate import validate  # noqa: E402

PROJECT_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = PROJECT_ROOT / "tests" / "fixtures" / "canvas"

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# Caches churn on their own during a test run, and a dot-directory in the project root can be an
# entire second checkout (a git worktree, say) -- walking one would make this test crawl for no
# added signal. Prune them; the renderer would never write there anyway.
PRUNED_DIRS = {"__pycache__", "node_modules"}


def load(name: str) -> dict:
    return json.loads((FIXTURES / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def client() -> TestClient:
    # `with` runs the lifespan -- which is where the two pin gates fire. If the template has drifted
    # this raises here, and it SHOULD: fail closed, never fall back to another layout.
    with TestClient(app) as c:
        yield c


def _tree(root: Path) -> set[str]:
    """Every file under `root`, pruning caches and dot-directories."""
    seen: set[str] = set()
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [
            d for d in dirnames if d not in PRUNED_DIRS and not d.startswith(".")
        ]
        for name in filenames:
            seen.add(str(Path(dirpath, name).relative_to(root)))
    return seen


# --- the editor is served, same-origin ---------------------------------------------------------


def test_index_is_served(client: TestClient):
    r = client.get("/")
    assert r.status_code == 200
    assert 'id="stage"' in r.text


def test_static_files_are_served(client: TestClient):
    assert client.get("/static/editor.js").status_code == 200
    assert client.get("/static/style.css").status_code == 200
    assert client.get("/static/interactions.js").status_code == 200


# --- /api/context: the editor NEVER hard-codes the SHA -----------------------------------------


def test_context_serves_the_pinned_sha(client: TestClient):
    """The whole reason /api/context exists.

    A hard-coded SHA in editor.js would turn re-pinning the template -- a deliberate, reviewed act
    (Sec 4.3) -- into a silent 409 for every user, whose only cure is a hard refresh nobody would
    guess. Serving it keeps reference/template_pin.json the single source of truth on both sides.
    """
    ctx = client.get("/api/context").json()
    assert ctx["template_sha256"] == load_pin()["renderer_template"]["sha256"]
    assert ctx["layout_index"] == 3
    assert ctx["canvas"] == {"w": 960, "h": 540, "unit": "pt"}
    assert ctx["brand"]["forbidden"] == "#4472C4"
    assert ctx["brand"]["red"] == "#FF0032"


def test_editor_js_does_not_hard_code_a_sha(client: TestClient):
    """Belt and braces with Gate 1: the SHA must arrive over the wire, not in the bundle."""
    sha = load_pin()["renderer_template"]["sha256"]
    assert sha not in client.get("/static/editor.js").text


# --- the happy path ----------------------------------------------------------------------------


def test_export_minimal_returns_a_one_slide_pptx(client: TestClient):
    r = client.post("/api/export", json=load("minimal"))

    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == PPTX_MIME
    assert r.content[:2] == b"PK"  # a .pptx is a zip, and a zip starts PK

    prs = Presentation(BytesIO(r.content))
    assert len(prs.slides) == 1


def test_export_sets_an_rfc5987_filename(client: TestClient):
    """The title is Korean often enough that the ASCII-only form is simply wrong."""
    ir = load("minimal")
    ir["title"] = "분기 리뷰"
    r = client.post("/api/export", json=ir)

    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert cd.startswith("attachment;")
    assert "filename*=UTF-8''" in cd
    # Korean folds to nothing under ASCII, so the fallback degrades to a safe stem rather than to
    # an empty or a mangled one.
    assert 'filename="deck.pptx"' in cd


def test_export_renders_connectors(client: TestClient):
    r = client.post("/api/export", json=load("connectors"))
    assert r.status_code == 200
    assert len(Presentation(BytesIO(r.content)).slides) == 1


# --- the error mapping -------------------------------------------------------------------------


def test_reject_header_is_422(client: TestClient):
    r = client.post("/api/export", json=load("reject_header"))
    assert r.status_code == 422
    body = r.json()
    assert body["error"] == "validation"
    assert "HEADER_RESERVED" in body["detail"]


def test_reject_accent1_is_422(client: TestClient):
    r = client.post("/api/export", json=load("reject_accent1"))
    assert r.status_code == 422
    assert "4472C4" in r.json()["detail"]


def test_stale_sha_is_409(client: TestClient):
    """A wrong template SHA is STALE, not invalid. The editor must say "reload", not "fix your deck".

    64 hex digits, so it passes the schema pattern and reaches the pin comparison in
    render_canvas.render() -- which is the code path under test.
    """
    ir = load("minimal")
    ir["target"]["template_sha256"] = "0" * 64
    r = client.post("/api/export", json=ir)

    assert r.status_code == 409, (
        f"a stale SHA came back as {r.status_code}. StalePayloadError SUBCLASSES "
        f"CanvasValidationError and must be caught BEFORE it, or every stale payload reports as a "
        f"422 and the user is told to fix a deck that is fine."
    )
    assert r.json()["error"] == "stale"


def test_stale_canvas_is_409_and_is_never_rescaled(client: TestClient):
    ir = load("minimal")
    ir["canvas"] = {"w": 1920, "h": 1080, "unit": "pt"}
    r = client.post("/api/export", json=ir)

    assert r.status_code == 409
    body = r.json()
    assert body["error"] == "stale"
    assert "never rescaled" in body["detail"]


def test_malformed_json_is_400_and_leaks_no_traceback(client: TestClient):
    r = client.post(
        "/api/export",
        content=b'{"schema": "canvas-ir/1", ',
        headers={"Content-Type": "application/json"},
    )
    assert r.status_code in (400, 422)
    assert "Traceback" not in r.text
    assert "File \"" not in r.text


def test_an_oversized_body_is_refused(client: TestClient):
    from aippt.canvas.app import MAX_BODY_BYTES

    r = client.post(
        "/api/export",
        content=b"x" * (MAX_BODY_BYTES + 1),
        headers={"Content-Type": "application/json"},
    )
    assert r.status_code == 413


# --- the security posture ----------------------------------------------------------------------


def test_the_template_path_in_the_payload_is_ignored(client: TestClient):
    """`target.template` is INFORMATIONAL. Honouring it would be an arbitrary-file read.

    A payload pointing at some other file must render exactly as the honest one does -- the server
    uses its own pinned path and never looks at this field.
    """
    ir = load("minimal")
    ir["target"]["template"] = "../../../../etc/passwd"
    r = client.post("/api/export", json=ir)

    assert r.status_code == 200
    assert len(Presentation(BytesIO(r.content)).slides) == 1


def test_no_cors_middleware_is_installed():
    """Same-origin only. A CORS grant would only widen who can reach an unauthenticated renderer."""
    names = [m.cls.__name__ for m in app.user_middleware]
    assert "CORSMiddleware" not in names


def test_the_export_writes_nothing_to_disk(client: TestClient):
    """prs.save() goes to a BytesIO. No temp file, no output directory, nothing on disk."""
    before = _tree(PROJECT_ROOT)
    r = client.post("/api/export", json=load("minimal"))
    assert r.status_code == 200
    assert _tree(PROJECT_ROOT) == before


# ===============================================================================================
# GATE 4b -- THE JSON IMPORT GOES THROUGH THE SAME GATE AS THE MOUSE
# ===============================================================================================
#
# Requirement 4 says the reserved regions are STRICTLY ENFORCED IN THE JS. They were enforced
# against the mouse (applyConstraints) and NOT against a file (loadIR): an element sitting on the
# Vieworks logo could be imported, rendered, and reserialized, and it came back as a 422 the user
# could not explain -- from a file this editor may itself have written before someone hand-edited
# it. These tests are what make "there is exactly one door into the stage, and it is locked" a
# checkable sentence instead of a claim in a comment.
#
# WHY THIS RUNS THE JS INSTEAD OF GREPPING IT
# -------------------------------------------
# Gate 1 (test_editor_source.py) is deliberately a grep, and says so. A grep cannot tell you that an
# element on the header is actually REFUSED -- only that some words appear in the file. So this
# section EXECUTES the shipped rule. It does not re-implement it in Python: a Python copy of
# illegalReason() would be a third source of truth, and the entire bug being fixed here is what
# happens when the rule exists in more than one place.
#
# The functions below are lifted VERBATIM out of editor.js by name and run in Node. Only the DOM
# around them is stubbed (stage, makeEl, titleField...), because the DOM is not what is under test:
# the judge and loadIR()'s ordering are. If node is not installed these tests SKIP -- `pytest tests/`
# must stay green for someone who installed only the renderer -- so the never-skipped structural
# assertions at the bottom hold the same invariants a second way.

EDITOR_JS = PROJECT_ROOT / "src" / "aippt" / "canvas" / "static" / "editor.js"
EDITOR_SRC = EDITOR_JS.read_text(encoding="utf-8")

NODE = shutil.which("node")

needs_node = pytest.mark.skipif(
    NODE is None,
    reason="node is not on PATH; the import gate's BEHAVIOUR cannot be executed. The structural "
    "assertions in this file still run, and Gate 1 still runs.",
)

# Every declaration and function the import path needs, in dependency order. Lifting by name (rather
# than slicing by line number) is what stops this harness from silently testing the wrong thing when
# editor.js is edited: a rename fails loudly here instead of quietly extracting nothing.
_DECLS = [
    "CANVAS",
    "SCHEMA",
    "HEADER_RESERVED",
    "FOOTER_RESERVED",
    "REJECT_REGIONS",
    "EPS",
    "BRAND_RED",
    "SHAPE_TYPES",
    "DEFAULTS",
    "GEOM_FIELDS",
    "ELEMENT_TYPES",
]

_FUNCS = [
    "intersects",
    "elementBBox",
    "outOfBounds",
    "esc",
    "geometryIsReadable",
    "illegalReason",
    "stateFromIR",
    "paragraphsToHTML",
    "loadIR",
]


def _js_decl(name: str) -> str:
    """The text of `var <name> = ...;` -- up to the first `;` at bracket depth 0."""
    start = EDITOR_SRC.index(f"var {name} = ")
    depth = 0
    for i in range(start, len(EDITOR_SRC)):
        ch = EDITOR_SRC[i]
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth -= 1
        elif ch == ";" and depth == 0:
            return EDITOR_SRC[start : i + 1]
    raise AssertionError(f"editor.js: `var {name} = ...` is never terminated")


def _js_function(name: str) -> str:
    """The text of `function <name>(...) {...}`, braces balanced.

    Same technique as Gate 1's _function_body, and it carries the same (satisfied) assumption: none
    of the extracted functions contains a brace inside a string, a comment or a regex literal.
    """
    start = EDITOR_SRC.index(f"function {name}(")
    depth = 0
    for i in range(EDITOR_SRC.index("{", start), len(EDITOR_SRC)):
        if EDITOR_SRC[i] == "{":
            depth += 1
        elif EDITOR_SRC[i] == "}":
            depth -= 1
            if depth == 0:
                return EDITOR_SRC[start : i + 1]
    raise AssertionError(f"editor.js: unbalanced braces in {name}()")


def _js_code(name: str) -> str:
    """_js_function(), with the comments stripped out.

    An ORDERING assertion has to key off what the code DOES, never off prose that happens to name
    the thing the code must not do first. restore()'s preflight comment explains why stopEditing()
    may not run above it -- and that sentence, left in, would satisfy an index('stopEditing()')
    looking for the call itself, so the gate would pass on a restore() that had been broken again.

    Carries the same (satisfied) assumption as _js_function: no `//` or `/*` inside a string literal
    in any extracted function.
    """
    body = re.sub(r"/\*.*?\*/", "", _js_function(name), flags=re.S)
    return re.sub(r"//[^\n]*", "", body)


# The DOM, stubbed. loadIR() is allowed to touch exactly these, and every one of them RECORDS what
# was done to it -- which is how "the stage was never touched" becomes an assertion and not a hope.
_STUBS = """
'use strict';
var madeIds = [];
var appended = 0;
var removed = 0;
var bodies = [];
var children = [{ remove: function () { removed += 1; } }];   // the deck the user already had
var titleField = { textContent: 'THE DECK THE USER ALREADY HAD' };
var statusLine = '';
var stage = {
  querySelectorAll: function () { return children; },
  appendChild: function () { appended += 1; },
};
function select() {}
function setStatus(m) { statusLine = m; }
function makeEl(s) {
  madeIds.push(s.id);
  var body = { innerHTML: '' };
  bodies.push(body);
  return { querySelector: function () { return body; } };
}
"""

_MAIN = """
var out = { threw: false, error: '' };
try {
  loadIR(JSON.parse(require('fs').readFileSync(0, 'utf8')));
} catch (e) {
  out.threw = true;
  out.error = String((e && e.message) || e);
}
out.made = madeIds;
out.appended = appended;
out.removed = removed;
out.title = titleField.textContent;
out.status = statusLine;
out.bodies = bodies.map(function (b) { return b.innerHTML; });
process.stdout.write(JSON.stringify(out));
"""


def _harness() -> str:
    parts = [_STUBS]
    parts += [_js_decl(n) for n in _DECLS]
    parts += [_js_function(n) for n in _FUNCS]
    parts.append(_MAIN)
    return "\n\n".join(parts)


def _import(ir: dict) -> dict:
    """Run the REAL loadIR() over `ir`. Returns what the import did to the stage."""
    with tempfile.TemporaryDirectory() as d:  # NOT under PROJECT_ROOT -- see _tree()
        harness = Path(d) / "import_gate.js"
        harness.write_text(_harness(), encoding="utf-8")
        proc = subprocess.run(
            [NODE, str(harness)],
            input=json.dumps(ir),
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=60,
        )
    assert proc.returncode == 0, (
        f"the harness itself failed to run -- this is a broken TEST, not a failing gate:\n"
        f"{proc.stderr}"
    )
    return json.loads(proc.stdout)


def deck(*elements: dict, title: str = "Q3 Review") -> dict:
    return {
        "schema": "canvas-ir/1",
        "canvas": {"w": 960, "h": 540, "unit": "pt"},
        "title": title,
        "elements": list(elements),
    }


def rect(**over) -> dict:
    """A legal rect in the body. Override or delete fields to make it illegal."""
    e = {"id": "el1", "type": "rect", "x": 100.0, "y": 150.0, "w": 120.0, "h": 80.0, "fill": "#FF0032"}
    e.update(over)
    return e


def _assert_stage_untouched(r: dict) -> None:
    """The whole of "reject the ENTIRE import". Nothing made, nothing cleared, title still there."""
    assert r["threw"], "the import was accepted"
    assert r["made"] == [], f"an element reached makeEl() anyway: {r['made']}"
    assert r["appended"] == 0, "an element was appended to the stage"
    assert r["removed"] == 0, (
        "the stage was CLEARED before the file was judged -- the user's deck is gone and the import "
        "failed anyway. This is the half-import the two-pass order exists to prevent."
    )
    assert r["title"] == "THE DECK THE USER ALREADY HAD", "the title was overwritten"


# --- the reserved regions, against a file --------------------------------------------------------


@needs_node
def test_import_refuses_an_element_in_the_header():
    """THE ONE THE CRITIC ASKED FOR. y 50..130 sits on the logo/divider band (0..92)."""
    r = _import(deck(rect(y=50.0)))
    assert "HEADER_RESERVED" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_an_element_in_the_footer():
    """x 800..900, y 505..535 sits on the slide-number chrome (733.6, 500 -> 960, 540)."""
    r = _import(deck(rect(x=800.0, y=505.0, w=100.0, h=30.0)))
    assert "FOOTER_RESERVED" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_a_rotated_element_whose_FOOTPRINT_enters_the_header():
    """The rotated-AABB half of the rule, which the path box alone would wave through.

    Unrotated this box is y 95..115 and legal -- it clears the header by 3 units. Turned 45 degrees
    about its centre its true footprint reaches y ~= 27, deep into the logo. The backend checks the
    footprint (geometry.py::element_bbox), so the editor must too, or the deck imports cleanly and
    exports as a 422.
    """
    r = _import(deck(rect(x=400.0, y=95.0, w=200.0, h=20.0, rotation=45)))
    assert "HEADER_RESERVED" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_a_connector_drawn_across_the_header():
    conn = {
        "id": "c1", "type": "connector", "kind": "straight",
        "x1": 100.0, "y1": 20.0, "x2": 300.0, "y2": 30.0,
        "color": "#FF0032", "width": 1.5, "arrow": "end",
    }
    r = _import(deck(conn))
    assert "HEADER_RESERVED" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_an_element_off_the_slide():
    r = _import(deck(rect(x=900.0, w=200.0)))  # x2 == 1100, past the right edge
    assert "outside" in r["error"], r["error"]
    _assert_stage_untouched(r)


# --- the NaN trap: a correct geometry check is not enough -----------------------------------------


@needs_node
def test_import_refuses_a_MISSING_coordinate():
    """THE TRAP. Fix the missing check without fixing this one and the gate still has a hole.

    A missing x becomes +undefined === NaN. EVERY comparison against NaN is false, so outOfBounds()
    returns false and intersects() returns false, and this element passes a PERFECTLY CORRECT
    reserved-region check. The judge has to prove the numbers are numbers before it compares them.
    """
    e = rect(y=50.0)  # in the header, too -- so a NaN-blind gate would admit a HEADER element
    del e["x"]
    r = _import(deck(e))
    assert "non-numeric" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_a_STRING_coordinate():
    """`"100"` is not a coordinate: the schema says `number`, so this is a 422 in waiting."""
    r = _import(deck(rect(x="100")))
    assert "non-numeric" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_a_BOOLEAN_rotation():
    """validate.py::_num rejects bool explicitly; typeof does the same job on this side."""
    r = _import(deck(rect(rotation=True)))
    assert "non-numeric" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_an_inverted_box():
    """A negative w is not a small box: elementBBox() would return x2 < x1 and every predicate
    downstream would then be comparing an interval that runs backwards."""
    r = _import(deck(rect(w=-120.0)))
    assert "non-positive" in r["error"], r["error"]
    _assert_stage_untouched(r)


# --- all or nothing -------------------------------------------------------------------------------


@needs_node
def test_ONE_illegal_element_refuses_the_WHOLE_import():
    """The critic's "reject the entire import", and the reason the stage is judged before it is
    cleared. The legal element must not land either."""
    r = _import(deck(rect(id="good", y=150.0), rect(id="bad", y=50.0)))
    assert "HEADER_RESERVED" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_the_refusal_names_EVERY_bad_element_not_just_the_first():
    """One pass, one message, numbered -- the same shape as the backend's 422 list."""
    r = _import(deck(rect(id="a", y=50.0), rect(id="b", y=505.0, x=800.0, w=100.0, h=30.0)))
    assert 'elements[0] (id "a", type "rect")' in r["error"], r["error"]
    assert 'elements[1] (id "b", type "rect")' in r["error"], r["error"]
    assert "HEADER_RESERVED" in r["error"] and "FOOTER_RESERVED" in r["error"]


@needs_node
def test_import_refuses_a_duplicate_id():
    r = _import(deck(rect(id="same"), rect(id="same", y=300.0)))
    assert "duplicate id" in r["error"], r["error"]
    _assert_stage_untouched(r)


@needs_node
def test_import_refuses_an_unknown_type():
    r = _import(deck(rect(type="dodecahedron")))
    assert "unknown type" in r["error"], r["error"]
    _assert_stage_untouched(r)


# --- and it still loads the decks it should -------------------------------------------------------
#
# A gate that refuses everything passes every test above and is useless. These are the other half.


@needs_node
def test_a_legal_deck_still_imports():
    r = _import(deck(rect(), title="Q3 Review"))
    assert not r["threw"], r["error"]
    assert r["made"] == ["el1"]
    assert r["appended"] == 1
    assert r["removed"] == 1, "the previous deck should have been cleared"
    assert r["title"] == "Q3 Review"
    assert "Loaded 1 element" in r["status"]


@needs_node
def test_an_element_flush_against_the_header_is_LEGAL():
    """HALF-OPEN, on this side too. BODY_SAFE starts at exactly 92, so a top-aligned element whose
    top edge IS 92 must import. Get this wrong and the gate spuriously refuses honest decks."""
    r = _import(deck(rect(y=92.0, h=100.0)))
    assert not r["threw"], r["error"]
    assert r["made"] == ["el1"]


@needs_node
def test_a_2_unit_element_imports_even_though_the_mouse_could_not_draw_one():
    """MIN_EXTENT (w >= 4) is drawing ergonomics; the LEGALITY rule is w > 0, which is what the
    schema says. The import borrows the judge and NOT the transform -- if it borrowed the transform
    it would silently resize the user's file to make it loadable."""
    r = _import(deck(rect(w=2.0, h=2.0)))
    assert not r["threw"], r["error"]
    assert r["made"] == ["el1"]


@needs_node
def test_a_paragraph_with_no_runs_does_not_crash_the_import():
    """`text.paragraphs[0].runs[0]` was unguarded, so a hand-edited paragraph was a TypeError -- and
    it was thrown AFTER the stage had been cleared, which cost the user the deck they had open."""
    e = rect(text={"align": "CENTER", "valign": "MIDDLE", "paragraphs": [{}]})
    r = _import(deck(e))
    assert not r["threw"], r["error"]
    assert r["made"] == ["el1"]


@needs_node
def test_a_run_with_no_text_does_not_crash_the_import():
    e = rect(text={"align": "CENTER", "valign": "MIDDLE", "paragraphs": [{"runs": [{}]}]})
    r = _import(deck(e))
    assert not r["threw"], r["error"]
    assert r["bodies"] == ["<div><span style=\"color: #FFFFFF\"></span></div>"], r["bodies"]


# --- the structural gates: these never skip, node or no node --------------------------------------


def _body(name: str) -> str:
    return _js_function(name)


def test_there_is_exactly_ONE_copy_of_the_legality_rule():
    """applyConstraints() must not carry its own reserved-region loop any more.

    Two copies of the rule is how the mouse and the importer drift apart -- and a drift here is a
    stage that looks legal and a 422 that cannot be explained.
    """
    body = _body("applyConstraints")
    assert "illegalReason(" in body, "applyConstraints() no longer calls the shared judge"
    assert "REJECT_REGIONS" not in body, (
        "applyConstraints() still has its own copy of the reserved-region check. The rule lives in "
        "illegalReason() and nowhere else."
    )


def test_the_import_path_calls_the_SAME_judge():
    assert "illegalReason(" in _body("loadIR"), (
        "loadIR() does not call illegalReason(). That is the entire defect: the file gets into the "
        "stage through a door the mouse is not allowed to use."
    )


def test_loadIR_judges_the_whole_file_BEFORE_it_touches_the_stage():
    """The ordering defect, machine-checked.

    Even a correct per-element check is worthless if it runs after the stage has been cleared: the
    first bad element then leaves a half-imported stage and the user's own deck already gone. So the
    judge, and the throw, must both come before ANY mutation.
    """
    body = _body("loadIR")

    judged = body.index("illegalReason(")
    thrown = body.index("if (problems.length)")

    mutations = {
        "clearing the stage": ".remove()",
        "overwriting the title": "titleField.textContent =",
        "building an element": "makeEl(",
        "appending to the stage": "appendChild(",
    }
    for what, needle in mutations.items():
        assert needle in body, f"loadIR() no longer does {what!r} -- this test is out of date"
        assert judged < body.index(needle), (
            f"loadIR() is {what} at a point BEFORE it has judged the elements. Reject-the-entire-"
            f"import is not reachable from there."
        )
        assert thrown < body.index(needle), (
            f"loadIR() is {what} before it raises on the collected problems, so a bad file would "
            f"still leave a half-imported stage behind."
        )


def test_makeEl_is_never_reached_without_a_verdict():
    """Every call site of makeEl() -- draw, duplicate, import, UNDO -- sits downstream of the judge."""
    for name in ("loadIR", "duplicateSelected"):
        body = _body(name)
        assert "makeEl(" in body
    # The draw path builds its element from applyConstraints()'s return value and bails on null.
    assert "if (!legal) {" in EDITOR_SRC, "the draw path no longer bails when the judge refuses"

    # restore() is the fourth call site and it WAS the bypass: it reached makeEl() with no verdict at
    # all, because "a restore is not an edit" was read as "a restore is not judged". Its verdict comes
    # from unrestorableReason(), which is illegalReason() applied to every record in the snapshot.
    body = _js_code("restore")
    assert body.index("unrestorableReason(") < body.index("makeEl("), (
        "restore() reaches makeEl() before it has a verdict -- undo is a door into the stage that "
        "the judge does not guard"
    )


# ===============================================================================================
# GATE 5 -- PHASE 2: THE PROPERTY PANEL, THE EDIT CONTRACT, AND UNDO
# ===============================================================================================
#
# Phase 2 relocated the property controls into a right-hand panel that appears on selection, made
# `contenteditable` an attribute that comes and goes, added run.bold / run.italic, and built the
# undo/redo stack Phase 1B deferred.
#
# Three of those four are cheap to get subtly, silently wrong, so each has a gate here:
#
#   - the panel   -- a control that is VISIBLE but patches nothing is the bug being fixed, and the
#                    CSS specificity trap below (`#props label` beats `[hidden]`) reintroduces it
#                    without touching a line of JS.
#   - the text    -- `contenteditable` was permanent, and only `pointer-events: none` kept the caret
#                    out. A class is not an attribute and editingText() consults the attribute.
#   - bold/italic -- an ELEMENT-level control over a PER-RUN property. The round trip is where this
#                    goes wrong: summarise the flag off the FIRST run and a deck with one bold word
#                    re-exports with the whole sentence bold.
#
# Undo is checked structurally (every mutation pushes) rather than behaviourally: exercising it needs
# a real DOM, and stubbing one convincingly enough to trust the result would be a bigger lie than the
# grep. What CAN be checked without a browser is checked, and it is checked here.

INDEX_HTML = (PROJECT_ROOT / "src" / "aippt" / "canvas" / "static" / "index.html").read_text(encoding="utf-8")
STYLE_CSS = (PROJECT_ROOT / "src" / "aippt" / "canvas" / "static" / "style.css").read_text(encoding="utf-8")

# Every styling property canvas-ir/1 actually HAS, and the control that drives it. If a control here
# stops existing the panel has lost a property; if one appears that is not here, it is patching
# something the schema cannot carry and validate.py will answer with a 422.
IR_PROPERTY_CONTROLS = {
    "p-fill": "fill",
    "p-stroke": "stroke.color / connector.color",
    "p-stroke-w": "stroke.width / connector.width",
    "p-color": "run.color",
    "p-size": "run.size",
    "p-bold": "run.bold",
    "p-italic": "run.italic",
    "p-align": "text.align",
    "p-valign": "text.valign",
    "p-kind": "connector.kind",
    "p-arrow": "connector.arrow",
    "p-rot": "rotation",
}


def _panel() -> str:
    """The markup of the right-hand property panel, and nothing else."""
    start = INDEX_HTML.index('<aside id="props"')
    return INDEX_HTML[start : INDEX_HTML.index("</aside>", start)]


def _toolbar() -> str:
    start = INDEX_HTML.index('<aside id="toolbar">')
    return INDEX_HTML[start : INDEX_HTML.index("</aside>", start)]


def _composer() -> str:
    start = INDEX_HTML.index('<section id="ai"')
    return INDEX_HTML[start : INDEX_HTML.index("</section>", start)]


# --- the panel ----------------------------------------------------------------------------------


def test_the_property_panel_exists_and_is_hidden_until_something_is_selected():
    """A panel of controls that patch nothing is a panel that lies.

    In Phase 1B these same inputs sat in the left toolbar and were permanently visible: with an empty
    selection every one of them was live and inert at the same time, and patchSelected() swallowed
    the edit.
    """
    assert re.search(r'<aside id="props"[^>]*\bhidden\b', INDEX_HTML), (
        "#props must ship with the `hidden` attribute -- the editor opens with nothing selected, so "
        "the panel's first frame must already be gone, not hidden a tick later by syncProps()."
    )
    assert "props.hidden = !s" in EDITOR_SRC, "syncProps() no longer shows/hides the panel"


@pytest.mark.parametrize("control", sorted(IR_PROPERTY_CONTROLS))
def test_every_ir_property_has_a_control_and_it_lives_in_the_panel(control: str):
    assert f'id="{control}"' in _panel(), (
        f"#{control} drives {IR_PROPERTY_CONTROLS[control]} and is not in the property panel."
    )


@pytest.mark.parametrize("control", sorted(IR_PROPERTY_CONTROLS))
def test_no_property_control_is_left_behind_in_the_toolbar(control: str):
    """The relocation has to be a MOVE. A copy left in the toolbar is a second control writing the
    same dataset field, and only one of them would be reflected by syncProps()."""
    assert f'id="{control}"' not in _toolbar()


def test_undo_and_redo_are_in_the_TOOLBAR_and_not_in_the_panel():
    """The one pair of controls that must work with NOTHING selected -- which is exactly the state
    the property panel is hidden in. Undo after a Delete is the whole point of undo, and after a
    Delete there is no selection."""
    assert 'id="b-undo"' in _toolbar() and 'id="b-redo"' in _toolbar()
    assert 'id="b-undo"' not in _panel() and 'id="b-redo"' not in _panel()


def test_the_hidden_attribute_actually_hides_inside_the_panel():
    """THE CSS SPECIFICITY TRAP, and it is not a style nit -- it is the whole scoping mechanism.

    `[hidden]` is a UA rule at (0,0,1,0). `#props label` is (0,1,0,1) and sets `display: flex`, so
    the ID SELECTOR WINS: a label carrying `hidden` would be announced as hidden, excluded from the
    accessibility tree -- and still be painted, and still be clickable. syncProps()'s data-for
    scoping would silently do nothing, and a connector would be offered a fill it cannot serialize.
    """
    assert "#props [hidden]" in STYLE_CSS, (
        "style.css does not override `display` for hidden descendants of #props, so `#props label` "
        "wins on specificity and the data-for scoping hides nothing."
    )
    assert "#props[hidden]" in STYLE_CSS, "the panel itself is never hidden by CSS"


def test_the_panel_scopes_its_controls_by_element_type():
    """The IR gives a connector no fill, no text and no rotation. Offering those controls anyway
    means offering a property that cannot be serialized -- and the JS guards (`if (isConnector())
    return`) would then eat the edit in silence, which is worse than not offering it."""
    panel = _panel()
    assert 'data-for="connector"' in panel and 'data-for="shape"' in panel
    assert "getAttribute('data-for')" in EDITOR_SRC, "syncProps() does not act on the data-for scoping"


# --- double-click to edit -----------------------------------------------------------------------


def test_contenteditable_is_not_hard_coded_onto_the_text_body():
    """THE DEFECT. `.body` carried contenteditable="true" AT ALL TIMES, and only
    `.txt { pointer-events: none }` stopped the caret -- which is a POINTER rule.

    So the element was, to the keyboard and to every assistive technology, permanently editable; and
    editingText() -- which reads .isContentEditable -- could not tell "the caret is in the text" from
    "the text merely exists". That is the same predicate that guards the Delete key.
    """
    # Asserted against the two things makeEl() could actually DO, rather than against the word:
    # the function's own comment explains why the attribute is absent, and a test that could not
    # survive being commented is a test about spelling.
    made = _js_function("makeEl")
    assert 'contenteditable="true"' not in made, (
        "makeEl() still hard-codes contenteditable on .body. The attribute must be ADDED on "
        "dblclick and REMOVED on exit, not merely masked with pointer-events."
    )
    assert "setAttribute('contenteditable'" not in made, (
        "makeEl() sets contenteditable on creation, which is the same defect by a longer route."
    )
    assert 'class="body" contenteditable' not in INDEX_HTML


def test_double_click_adds_the_attribute_and_leaving_removes_it():
    assert re.search(
        r"dblclick[\s\S]{0,1200}setAttribute\('contenteditable', 'true'\)", EDITOR_SRC
    ), "the dblclick handler does not turn .body into a contenteditable"
    assert "removeAttribute('contenteditable')" in _js_function("stopEditing"), (
        "stopEditing() drops the .editing class but leaves contenteditable behind -- which is the "
        "defect, just one class further along."
    )


def test_a_single_click_still_selects_and_drags():
    """The regression that would cost a whole phase. .txt must stay pointer-transparent until the
    element is being edited, or a no-fill textbox can never be grabbed at all."""
    assert re.search(r"\.el\s+\.txt\s*\{[^}]*pointer-events:\s*none", STYLE_CSS, re.S)
    assert re.search(r"\.el\.editing\s+\.txt\s*\{[^}]*pointer-events:\s*auto", STYLE_CSS, re.S)


def test_delete_and_backspace_still_remove_the_selection_and_still_not_while_typing():
    """Phase 1B already shipped this. It is here so that Phase 2 cannot quietly take it away."""
    assert "function deleteSelected()" in EDITOR_SRC
    assert re.search(
        r"if \(editingText\(ev\.target\)\) return;[\s\S]{0,300}"
        r"ev\.key === 'Delete' \|\| ev\.key === 'Backspace'",
        EDITOR_SRC,
    ), (
        "Delete/Backspace is no longer guarded by editingText(), so pressing Backspace while typing "
        "would delete the element you are typing into."
    )


# --- undo / redo --------------------------------------------------------------------------------


def test_undo_and_redo_exist_and_are_bound_to_the_keyboard():
    assert "function undo()" in EDITOR_SRC and "function redo()" in EDITOR_SRC
    assert re.search(r"k === 'z' \|\| k === 'y'", EDITOR_SRC), "Ctrl+Z / Ctrl+Y are not bound"
    assert "ev.shiftKey) redo()" in EDITOR_SRC, "Ctrl+Shift+Z does not redo"


def test_undo_is_wired_through_commit_the_one_door_state_changes_through():
    """commit() is the single place dataset geometry is written. Hanging history anywhere else means
    some mutation, somewhere, is not undoable -- and you find out which one from a user."""
    assert "noteChange(" in _js_function("commit")


@pytest.mark.parametrize("mutation", ["deleteSelected", "duplicateSelected"])
def test_the_mutations_that_bypass_commit_push_their_own_undo_step(mutation: str):
    """Delete and Duplicate do not go through commit() -- they add and remove NODES. They have to
    push for themselves, and they have to push BEFORE they mutate: a snapshot taken afterwards
    records the stage without the element, and Delete becomes permanent with extra steps."""
    assert "pushUndo(snapshot())" in _js_function(mutation)


def test_a_refused_frame_costs_no_undo_step():
    """Reject-and-hold means an illegal drag frame changes nothing. If it still pushed, dragging a
    box against the header would fill the stack with states the user cannot tell apart."""
    body = _js_function("commit")
    refused = body.index("flashBlocked(el); return false;")
    pushed = body.index("noteChange(")
    assert refused < pushed, "commit() pushes history before it knows the frame was accepted"


def test_undo_restores_through_makeEl_and_never_measures_the_dom():
    """A snapshot is JS STATE -- dataset, innerHTML, DOM order. Restoring one rebuilds the elements
    down the same one-way state -> CSS path every other write takes, so undo structurally CANNOT
    reintroduce a device pixel. (Gate 1 enforces the ban globally; this says WHERE it is kept.)"""
    body = _js_function("restore")
    assert "makeEl(" in body, "restore() no longer rebuilds elements from state"
    assert "getBoundingClientRect" not in body


# --- undo is the THIRD door into the stage, and it is judged like the other two --------------------
#
# The test above says restore() rebuilds from state and never measures. That is necessary and it is
# not sufficient: it passed happily while restore() was sending snapshot geometry straight through
# makeEl()/appendChild() with no call to illegalReason() at all. What follows is the missing half.
#
# unrestorableReason() is lifted out of editor.js by name and executed in Node, exactly like loadIR()
# above -- but with NO STUBS AT ALL. There is no `stage`, no `titleField`, no `makeEl` in that
# harness, so if the preflight ever reaches for the DOM it is a ReferenceError and the gate says so.
# "It touches no DOM" is thereby an assertion and not a hope, which is the whole reason the preflight
# is allowed to run before the stage has been cleared.

_PREFLIGHT_FUNCS = [
    "intersects",
    "elementBBox",
    "outOfBounds",
    "geometryIsReadable",
    "illegalReason",
    "unrestorableReason",
]

_PREFLIGHT_MAIN = """
var out = { threw: false, error: '', reason: null };
try {
  out.reason = unrestorableReason(JSON.parse(require('fs').readFileSync(0, 'utf8')));
} catch (e) {
  out.threw = true;
  out.error = String((e && e.message) || e);
}
process.stdout.write(JSON.stringify(out));
"""


def _preflight(snap) -> dict:
    """Run the REAL unrestorableReason() over `snap`, with no DOM in scope."""
    src = "\n\n".join(
        ["'use strict';"]
        + [_js_decl(n) for n in _DECLS]
        + [_js_function(n) for n in _PREFLIGHT_FUNCS]
        + [_PREFLIGHT_MAIN]
    )
    with tempfile.TemporaryDirectory() as d:  # NOT under PROJECT_ROOT -- see _tree()
        harness = Path(d) / "restore_gate.js"
        harness.write_text(src, encoding="utf-8")
        proc = subprocess.run(
            [NODE, str(harness)],
            input=json.dumps(snap),
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=60,
        )
    assert proc.returncode == 0, (
        f"the harness itself failed to run -- this is a broken TEST, not a failing gate:\n"
        f"{proc.stderr}"
    )
    out = json.loads(proc.stdout)
    assert not out["threw"], (
        "unrestorableReason() threw when handed a snapshot. It is called BEFORE the stage is "
        f"cleared, so it may refuse, but it may never throw and it may never touch the DOM: "
        f"{out['error']}"
    )
    return out


def el_state(**over) -> dict:
    """One element of a snapshot, shaped exactly as read() hands it back: geometry already NUMBERS."""
    s = {"id": "el1", "type": "rect", "x": 100.0, "y": 150.0, "w": 120.0, "h": 80.0, "rot": 0}
    s.update(over)
    return s


def snap(*states: dict, title: str = "Q3 Review") -> dict:
    return {
        "title": title,
        "els": [{"state": s, "html": None, "sel": False} for s in states],
    }


@needs_node
def test_a_snapshot_off_a_legal_stage_still_restores():
    """The judge is inert, so judging a restore must cost the user nothing. If this fails, undo is
    BROKEN -- every legal snapshot is being refused -- which is the one way this repair could be
    worse than the bug it fixes."""
    out = _preflight(snap(el_state(), el_state(id="el2", x=400.0, type="ellipse")))
    assert out["reason"] is None, f"a legal snapshot was refused: {out['reason']}"


@needs_node
def test_the_undo_preflight_refuses_a_snapshot_that_lands_on_a_reserved_region():
    """THE BUG. An element on the logo cannot be drawn with the mouse and cannot be imported from a
    file -- and, until now, could be handed back to the stage by pressing Ctrl+Z, whereupon the deck
    the editor swore was legal earns a 422 from the backend."""
    out = _preflight(snap(el_state(y=10.0, h=40.0)))
    assert out["reason"] is not None, "undo would put an element back on the header"
    assert "els[0]" in out["reason"] and "el1" in out["reason"]


@needs_node
@pytest.mark.parametrize(
    "bad",
    [
        pytest.param(el_state(x=None), id="missing-coordinate"),
        pytest.param(el_state(w="120"), id="string-coordinate"),
        pytest.param(el_state(w=0.0), id="zero-width"),
        pytest.param(el_state(h=-80.0), id="inverted-box"),
        pytest.param(el_state(x=900.0, w=400.0), id="off-the-slide"),
        pytest.param(el_state(type="star"), id="unknown-type"),  # NOT triangle -- triangle is legal
    ],
)
def test_the_undo_preflight_refuses_every_way_a_record_can_be_illegal(bad: dict):
    """The same six ways loadIR() refuses a file, because it is the same judge -- reached through a
    different door. A NaN coordinate matters most: every comparison against NaN is false, so an
    unjudged garbage record sails through the geometry checks downstream."""
    assert _preflight(snap(bad))["reason"] is not None


@needs_node
@pytest.mark.parametrize(
    "junk",
    [None, {}, {"title": "t"}, {"els": []}, {"title": "t", "els": "not a list"}],
    ids=["none", "empty", "no-els", "no-title", "els-not-a-list"],
)
def test_the_undo_preflight_refuses_a_history_entry_that_is_not_a_snapshot(junk):
    """restore() used to be handed whatever pop() returned. pop() on an empty array returns undefined
    -- and `undefined.els` is a TypeError thrown AFTER the stage was cleared."""
    assert _preflight(junk)["reason"] is not None


@needs_node
def test_the_undo_preflight_refuses_a_malformed_record_instead_of_throwing_on_it():
    out = _preflight({"title": "t", "els": [{"html": None, "sel": False}]})  # a record with no state
    assert out["reason"] is not None


# --- the same invariants, structurally: these run even where node does not -------------------------


def test_restore_judges_the_snapshot_before_it_touches_the_dom():
    """ORDER IS THE FIX. A check bolted into the loop cannot produce "abort the restoration" -- by
    the time that loop runs, the user's deck is already gone. And stopEditing() MUTATES THE DOM (it
    drops .editing and removes contenteditable) and ends the transaction, so "before touching the
    DOM" means above stopEditing(), not merely above the clear."""
    body = _js_code("restore")
    judged = body.index("unrestorableReason(")
    for dom_touch in ("stopEditing()", "n.remove()", "titleField.innerHTML", "makeEl("):
        assert judged < body.index(dom_touch), (
            f"restore() reaches {dom_touch} before it has judged the snapshot, so a refusal can only "
            f"leave a half-restored stage behind and the user's own deck already gone"
        )


def test_restore_reports_whether_it_actually_restored():
    """All-or-nothing is only all-or-nothing if the caller can tell which. A restore() that refuses
    silently and returns undefined lets undo() pop the stack anyway -- the data-loss path, rebuilt."""
    body = _js_code("restore")
    assert "return false;" in body and "return true;" in body


def test_the_undo_preflight_judges_and_transforms_nothing():
    """It is called before the stage is cleared, so it must be inert in both directions: it may not
    mutate the snapshot (snapping or clamping a restore would silently rewrite the user's own history
    under them -- the push-out we refused the mouse) and it may not touch the DOM."""
    body = _js_code("unrestorableReason")
    assert "illegalReason(" in body, "the undo preflight does not call the judge at all"
    for transform in ("applyConstraints(", "snapVal(", "snapAxis(", "clamp("):
        assert transform not in body, f"the undo preflight calls {transform} -- a judge that repairs"
    for dom in ("stage.", "titleField", "document.", "makeEl(", "getBoundingClientRect"):
        assert dom not in body, f"the undo preflight touches the DOM ({dom})"


@pytest.mark.parametrize("fn,stack", [("undo", "past"), ("redo", "future")])
def test_undo_and_redo_do_not_pop_until_the_restore_has_actually_succeeded(fn: str, stack: str):
    """`restore(past.pop())` evaluates the pop AS THE ARGUMENT -- the entry is off the stack before
    restore() has done anything. If the restore is then refused or throws, the stage is cleared AND
    the snapshot that would put it back is gone: the deck and the undo step, in one keystroke."""
    body = _js_code(fn)
    assert f"restore({stack}.pop())" not in body, (
        f"{fn}() pops as the argument, so a refused restore destroys the entry it refused"
    )
    restored = body.index("restore(")
    popped = body.index(f"{stack}.pop()")
    assert restored < popped, f"{fn}() still takes the entry off the stack before restoring it"


def test_undo_and_redo_bail_out_when_the_restore_was_refused():
    """On refusal nothing moved, so nothing may be pushed onto the other stack either -- a redo entry
    for a restore that never happened is a step that undoes something the user never saw."""
    for fn, other in (("undo", "future"), ("redo", "past")):
        body = _js_code(fn)
        assert "if (!restore(" in body and "return;" in body, f"{fn}() ignores restore()'s refusal"
        assert body.index("if (!restore(") < body.index(f"{other}.push("), (
            f"{fn}() pushes onto {other} before it knows the restore was accepted"
        )


def test_loadIR_is_not_where_the_import_pushes_its_undo_step():
    """loadIR() is lifted out of editor.js BY NAME and executed in Node by the Gate 4b harness above,
    against a stubbed stage. A call to snapshot() or pushUndo() from inside it is a ReferenceError in
    that harness -- the gate would fail on code that is perfectly correct in a browser. The import is
    one undo step either way; the CALLER is where it is pushed, and the caller is also the only one
    that knows whether loadIR() actually accepted the file."""
    body = _js_function("loadIR")
    assert "pushUndo" not in body and "snapshot(" not in body


def test_zoom_is_not_undoable():
    """Zoom is VIEW, not document. Undoing a move and having the zoom jump under you is not undo,
    it is time travel -- so the zoom slider is the one control NOT wired into a transaction, and the
    zoom is the one piece of editor state NOT in a snapshot."""
    assert "p-zoom" not in _js_function("snapshot")

    # The list of controls txControl() is applied to, and nothing else.
    wired = EDITOR_SRC[EDITOR_SRC.index("function txControl(") :]
    wired = wired[: wired.index("guardColor($('p-fill')")]
    assert "p-zoom" not in wired, "the zoom slider pushes undo steps"


# --- run.bold / run.italic: the round trip ------------------------------------------------------

_P2_DECLS = ["BRAND_RED", "DEFAULTS", "BLOCK_TAGS"]
_P2_FUNCS = ["stripC0", "hex2", "cssColorToHex", "colorOf", "normalizeText", "stateFromIR"]

# A DOM small enough to be obviously correct: normalizeText() reads nodeType, tagName, nodeValue,
# childNodes, lastChild, style.color and getAttribute('color'), and nothing else. It MEASURES
# nothing, which is exactly why a stub this thin is honest here and would not be for undo.
_P2_DOM = """
function T(t) { return { nodeType: 3, nodeValue: t }; }
function E(tag, kids) {
  var e = {
    nodeType: 1, tagName: tag, childNodes: kids || [], style: {},
    getAttribute: function () { return null; },
  };
  e.lastChild = e.childNodes.length ? e.childNodes[e.childNodes.length - 1] : null;
  return e;
}
"""


def _run_p2(main: str) -> object:
    parts = ["'use strict';", _P2_DOM]
    parts += [_js_decl(n) for n in _P2_DECLS]
    parts += [_js_function(n) for n in _P2_FUNCS]
    parts.append(main)
    with tempfile.TemporaryDirectory() as d:  # NOT under PROJECT_ROOT -- see _tree()
        harness = Path(d) / "p2_gate.js"
        harness.write_text("\n\n".join(parts), encoding="utf-8")
        proc = subprocess.run(
            [NODE, str(harness)], capture_output=True, text=True, encoding="utf-8", timeout=60
        )
    assert proc.returncode == 0, (
        f"the harness itself failed to run -- this is a broken TEST, not a failing gate:\n"
        f"{proc.stderr}"
    )
    return json.loads(proc.stdout)


@needs_node
def test_the_element_bold_flag_folds_into_every_run():
    """The IR has no element-level bold. run.bold is the only place the flag can land, so the
    serializer OR-s it into every run -- the same relationship size and colour already have."""
    out = _run_p2(
        "process.stdout.write(JSON.stringify("
        "normalizeText(E('DIV', [T('hi')]), 14, '#000000', true, false)));"
    )
    assert out == [{"runs": [{"t": "hi", "size": 14, "color": "#000000", "bold": True}]}], out


@needs_node
def test_a_b_tag_inside_the_body_still_bolds_its_own_run_with_the_flag_off():
    """NO REGRESSION. Per-run bold already worked -- it comes from <b>/<strong> in the
    contenteditable, which is where a paste and the browser's own Ctrl+B put it. The element flag is
    additive to that, never a replacement for it."""
    out = _run_p2(
        "process.stdout.write(JSON.stringify("
        "normalizeText(E('DIV', [E('B', [T('x')])]), 14, '#000000', false, false)));"
    )
    assert out[0]["runs"][0]["bold"] is True, out


@needs_node
def test_a_plain_run_stays_plain_when_the_flag_is_off():
    out = _run_p2(
        "process.stdout.write(JSON.stringify("
        "normalizeText(E('DIV', [T('hi')]), 14, '#000000', false, false)));"
    )
    assert "bold" not in out[0]["runs"][0], out
    assert "italic" not in out[0]["runs"][0], out


def _state_of(runs: list) -> dict:
    e = rect(text={"align": "CENTER", "valign": "MIDDLE", "paragraphs": [{"runs": runs}]})
    return _run_p2(
        f"process.stdout.write(JSON.stringify(stateFromIR({json.dumps(e)})));"
    )


@needs_node
def test_the_element_flag_is_a_SUMMARY_of_the_runs_and_not_the_first_one():
    """THE ROUND-TRIP TRAP, and the reason stateFromIR() counts instead of peeking.

    size and colour are read off the FIRST run. Reading bold the same way would be actively
    destructive: the flag is OR-ed back into EVERY run on the way out, so a deck with one bold word
    in a plain sentence would reload with the flag on and re-export with the WHOLE SENTENCE bold. A
    round trip must not change the deck.

    So mixed runs summarise to OFF, and their own <b>/<i> tags -- which paragraphsToHTML() writes --
    carry the per-run bold instead. Nothing is lost either way.
    """
    plain = {"t": "b", "size": 14, "color": "#FFFFFF"}
    bold = {"t": "a", "size": 14, "color": "#FFFFFF", "bold": True}

    assert _state_of([bold, dict(bold, t="c")])["bold"] is True, "every run bold -> flag ON"
    assert _state_of([bold, plain])["bold"] is False, (
        "the FIRST run is bold and the second is not. If the flag came back ON, re-exporting this "
        "deck would bold the plain run too -- the round trip would have rewritten the user's deck."
    )
    assert _state_of([plain, bold])["bold"] is False


@needs_node
def test_italic_summarises_the_same_way():
    italic = {"t": "a", "size": 14, "color": "#FFFFFF", "italic": True}
    assert _state_of([italic])["italic"] is True
    assert _state_of([italic, {"t": "b", "size": 14, "color": "#FFFFFF"}])["italic"] is False


@needs_node
def test_a_textless_element_is_neither_bold_nor_italic():
    """`nRuns > 0` is doing real work: with no runs at all, `nBold === nRuns` is `0 === 0`, and a
    plain rect would come back BOLD."""
    out = _run_p2(f"process.stdout.write(JSON.stringify(stateFromIR({json.dumps(rect())})));")
    assert out["bold"] is False and out["italic"] is False, out


@needs_node
def test_turning_bold_off_also_strips_the_tags():
    """The flag alone cannot turn bold OFF, and that asymmetry is the whole of setTextFlag().

    run.bold comes from <b>/<strong> in the body as well as from the flag. An "off" that merely
    stopped ADDING bold would leave the tags -- and therefore the bold -- in the exported deck, with
    the button showing un-pressed. The control would look broken because it would BE broken.
    """
    body = _js_function("setTextFlag")
    assert "unwrapTags(" in body, "setTextFlag() does not remove the tags when it clears the flag"
    assert "['b', 'strong']" in body and "['i', 'em']" in body, (
        "<strong> and <em> are what a paste and some editors produce; unwrapping only <b>/<i> "
        "leaves half the cases bold."
    )


# ===============================================================================================
# PHASE 3 -- /api/generate: THE MOCK AUTHOR, AND THE SEAM A REAL ONE SLOTS INTO
# ===============================================================================================
#
# The endpoint is a mock, and every test here is written to survive it ceasing to be one. What is
# under test is the CONTRACT: the payload is a valid canvas-ir/1 deck, it carries the LIVE pin's SHA
# rather than a literal, it survives the editor's own gate, and it costs what a model call would.
# A real author replaces mock_ir() and nothing else -- if that swap breaks anything below, it broke
# the contract, not the test.
#
# THE ONE THAT MATTERS MOST is test_the_generated_deck_EXPORTS: it feeds what /api/generate said
# straight into /api/export. That closes the loop through the real validator, the real pin check and
# the real renderer, and it is the only assertion here a merely plausible-looking payload could not
# satisfy.

APP_PY = (PROJECT_ROOT / "src" / "aippt" / "canvas" / "app.py").read_text(encoding="utf-8")

# "Do not use external APIs", checked against EVERY import line in app.py -- including the lazy ones
# (uvicorn is imported inside main()). A grep is crude and it is exactly right: there is no
# legitimate reason for any of these names to appear, and the day one does, this fails.
NETWORK_LIBS = ["httpx", "aiohttp", "openai", "anthropic", "urllib.request", "socket", "requests"]


@pytest.fixture(scope="module")
def generated(client: TestClient) -> dict:
    """ONE call, reused: every assertion below is about the same payload, and it costs 1.5 s."""
    r = client.post("/api/generate", json={"prompt": "Q3 마케팅 전략을 3단계 퍼널로 요약해줘"})
    assert r.status_code == 200, r.text
    return r.json()


def test_generate_returns_a_valid_canvas_ir_payload(generated: dict):
    validate(generated)  # the real validator, not a re-implementation of it
    assert generated["schema"] == "canvas-ir/1"
    assert generated["canvas"] == {"w": 960, "h": 540, "unit": "pt"}
    assert generated["target"]["layout_index"] == 3


def test_the_generated_deck_EXPORTS(client: TestClient, generated: dict):
    """The whole point. What the mock author says, the real renderer has to be able to draw."""
    r = client.post("/api/export", json=generated)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == PPTX_MIME
    assert len(Presentation(BytesIO(r.content)).slides) == 1


def test_the_mock_carries_the_LIVE_sha_and_not_a_literal(generated: dict):
    """A literal SHA in app.py goes stale the day the template is re-pinned, and then hands every
    user a 409 on a deck they did not author. Same argument that put /api/context there."""
    sha = load_pin()["renderer_template"]["sha256"]
    assert generated["target"]["template_sha256"] == sha
    assert sha not in APP_PY, "app.py hard-codes the template SHA; it must come from the pin"


def test_the_mock_is_a_DIAGRAM_and_not_a_wall_of_text(generated: dict):
    """3-4 shapes and connecting arrows. A payload that does not exercise the IR proves only that
    JSON can be posted over HTTP."""
    els = generated["elements"]
    assert len([e for e in els if e["type"] != "connector"]) >= 3
    arrows = [e for e in els if e["type"] == "connector"]
    assert arrows, "the mock draws no connectors"
    assert all(a["arrow"] != "none" for a in arrows), "a connecting arrow with no arrowhead"
    assert els[0]["type"] == "connector", (
        "ARRAY ORDER IS Z-ORDER: the connectors must come first so the lines sit BEHIND the cards "
        "they join -- the convention the editor follows when it inserts one."
    )
    assert len({e["id"] for e in els}) == len(els), "duplicate ids -- ids become shape names"


def test_the_prompt_is_NOT_read(client: TestClient):
    """It is a mock and it must not look like anything else. Two prompts, one deck."""
    a = client.post("/api/generate", json={"prompt": "a"}).json()
    b = client.post("/api/generate", json={"prompt": "완전히 다른 요청"}).json()
    assert a == b


def test_generate_costs_what_a_model_call_would_cost(client: TestClient):
    assert MOCK_LATENCY_S == 1.5
    assert "asyncio.sleep(MOCK_LATENCY_S)" in APP_PY
    t0 = time.perf_counter()
    assert client.post("/api/generate", json={"prompt": "x"}).status_code == 200
    # asyncio.sleep never returns early; the slack is for timer granularity, not for correctness.
    assert time.perf_counter() - t0 >= 1.2


def test_a_bad_request_is_refused_IMMEDIATELY_and_never_pays_the_latency(client: TestClient):
    t0 = time.perf_counter()
    r = client.post("/api/generate", json={"prompt": "   "})
    elapsed = time.perf_counter() - t0
    assert r.status_code == 422
    assert r.json()["error"] == "validation"
    assert elapsed < 1.0, f"the refusal slept anyway ({elapsed:.2f}s); no model would have been called"


@pytest.mark.parametrize("body", [{}, {"prompt": ""}, {"prompt": 3}, {"prompt": None}, [], "x"])
def test_generate_refuses_every_shape_of_missing_prompt(client: TestClient, body):
    assert client.post("/api/generate", json=body).status_code == 422


def test_generate_shares_the_body_guard_with_export(client: TestClient):
    """One guard, two routes. A second copy is a second thing to forget to update."""
    from aippt.canvas.app import MAX_BODY_BYTES

    big = client.post(
        "/api/generate",
        content=b"x" * (MAX_BODY_BYTES + 1),
        headers={"Content-Type": "application/json"},
    )
    assert big.status_code == 413

    bad = client.post(
        "/api/generate", content=b'{"prompt": ', headers={"Content-Type": "application/json"}
    )
    assert bad.status_code == 400
    assert "Traceback" not in bad.text


def test_generate_writes_nothing_to_disk(client: TestClient):
    before = _tree(PROJECT_ROOT)
    assert client.post("/api/generate", json={"prompt": "x"}).status_code == 200
    assert _tree(PROJECT_ROOT) == before


@pytest.mark.parametrize("lib", NETWORK_LIBS)
def test_the_mock_author_reaches_no_external_api(lib: str):
    imports = [ln for ln in APP_PY.splitlines() if re.match(r"\s*(import|from)\s", ln)]
    assert not any(lib in ln for ln in imports), f"app.py imports {lib}; the mock calls nothing"


def test_mock_ir_is_a_pure_function_of_the_pin():
    """The seam. A real author replaces THIS function -- so it takes the pin and returns a deck, with
    no request, no app state and no I/O of its own."""
    ir = mock_ir(load_pin())
    validate(ir)
    assert ir["target"]["template_sha256"] == load_pin()["renderer_template"]["sha256"]


# --- the editor's half of the seam ---------------------------------------------------------------


def test_the_prompt_and_the_generate_button_live_in_the_composer():
    comp = _composer()
    assert 'id="p-prompt"' in comp and 'id="b-generate"' in comp
    assert 'id="p-prompt"' not in _panel(), (
        "the prompt is not a property of the SELECTION. The property panel is hidden whenever "
        "nothing is selected -- which is exactly the state you write a prompt in."
    )


def test_generate_posts_the_prompt_to_the_mock_endpoint():
    body = _js_function("generate")
    assert "'/api/generate'" in body
    assert "JSON.stringify({ prompt: prompt })" in body
    assert "$('b-generate').addEventListener('click', generate)" in EDITOR_SRC


def test_generate_hands_the_payload_to_the_SAME_gate_the_file_import_uses():
    """No second copy of the judge. A deck a model wrote is no more trusted than a file a human
    hand-edited -- and it is the author most likely to put a box on the Vieworks logo."""
    body = _js_function("generate")
    assert "loadIR(ir)" in body
    assert "illegalReason" not in body, "generate() re-implements the judge; loadIR() IS the judge"


def test_generate_is_one_undo_step_and_the_CALLER_pushes_it():
    """It replaces the whole deck. loadIR() must not push for itself: it is lifted out by name and
    run in Node against a stubbed stage, where snapshot() does not exist."""
    body = _js_function("generate")
    assert "snapshot()" in body and "pushUndo(before)" in body
    assert "pushUndo" not in _js_function("loadIR")


def test_a_refused_draft_costs_no_undo_step():
    body = _js_code("generate")
    assert body.index("loadIR(ir)") < body.index("pushUndo(before)"), (
        "generate() pushes the undo step before it knows loadIR() accepted the draft"
    )


def test_the_loading_state_is_always_cleared():
    """The one thing worse than a slow button is a button that never comes back."""
    body = _js_code("generate")
    assert "setBusy(true)" in body and "finally" in body
    busy = _js_function("setBusy")
    assert "disabled = on" in busy, "the button is not disabled -- a double press posts twice"
    assert "aria-busy" in busy, "the loading state is invisible to a screen reader"


def test_generate_is_disabled_when_the_server_is_unreachable():
    assert "b-generate" in _js_function("disableExports"), (
        "the draft comes from the same server as the export; leaving the button live offers an "
        "action whose only possible outcome is the banner the user is already reading."
    )


# ===============================================================================================
# PHASE 4 -- /api/generate: THE LIVE CLIWORKERCLIENT AUTHOR, MOCKED AT THE NETWORK SEAM
# ===============================================================================================
#
# The mock stays the default and every Phase 3 test above still describes it. What is added here is
# the LIVE path with the one seam that shells out to a claude/codex CLI --
# CliWorkerClient.complete_json(), aippt.adapters.ai_adapter -- monkeypatched: the schema relaxation,
# the prompt injection, the pin injection and the validation around it run for real, and only the CLI
# round trip is stubbed. The mode is read at REQUEST time from AIPPT_LIVE_MODE, so setting the env
# var per-test flips it without a restart.

import aippt.adapters.ai_adapter as ai_adapter  # noqa: E402

# A deck shaped exactly as the system prompt asks the model to answer: schema + canvas fixed, NO
# target (the server injects it from the pin), one titled rect. It validates once target is
# injected. CliWorkerClient.complete_json() ALREADY returns a schema-checked dict -- JSON extraction
# and schema validation both happen inside ai_adapter.py -- so this is what the stub hands back
# directly, never something the test JSON-encodes first.
_LIVE_DECK = {
    "schema": "canvas-ir/1",
    "canvas": {"w": 960, "h": 540, "unit": "pt"},
    "title": "라이브 데크",
    "elements": [
        {
            "id": "box1", "type": "rect",
            "x": 100.0, "y": 150.0, "w": 220.0, "h": 110.0,
            "fill": "#E63312",
            "text": {
                "align": "CENTER", "valign": "MIDDLE",
                "paragraphs": [{"runs": [{"t": "실적 요약", "size": 18, "bold": True, "color": "#FFFFFF"}]}],
            },
        }
    ],
}


def _live_deck() -> dict:
    """A FRESH, independent copy of _LIVE_DECK. _finalize_live_ir() mutates nested dicts in place
    (filling in defaulted colors/arrows) -- handing out the module-level fixture itself, or a mere
    dict(_LIVE_DECK) shallow copy, would let one test's call leak mutations into every other test
    that reuses it."""
    return json.loads(json.dumps(_LIVE_DECK))


def _stub_worker(monkeypatch, returns=None, raises=None):
    def fake_complete_json(self, *, system, user, schema, max_tokens=8000):
        assert isinstance(user, str) and user.strip(), "the prompt reached the model empty"
        assert isinstance(system, str) and system.strip(), "no system prompt reached the model"
        assert "target" not in schema.get("required", []), (
            "the LIVE schema must not require 'target' -- the model cannot know the pinned SHA"
        )
        if raises is not None:
            raise raises
        return returns
    monkeypatch.setattr(ai_adapter.CliWorkerClient, "complete_json", fake_complete_json)


def test_context_reports_live_mode_when_the_env_var_is_set(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    ctx = client.get("/api/context").json()
    assert ctx["mode"] == "live"
    assert ctx["model"]  # a human string is surfaced for the badge tooltip


def test_context_reports_mock_mode_without_the_env_var(client, monkeypatch):
    monkeypatch.delenv("AIPPT_LIVE_MODE", raising=False)
    ctx = client.get("/api/context").json()
    assert ctx["mode"] == "mock"
    assert ctx["model"] is None


@pytest.mark.parametrize("value", ["", "0", "false", "True", "  ", "11", "yes"])
def test_only_the_literal_1_enables_live_mode(client, monkeypatch, value):
    monkeypatch.setenv("AIPPT_LIVE_MODE", value)
    assert client.get("/api/context").json()["mode"] == "mock"


def test_a_padded_1_still_enables_live_mode(client, monkeypatch):
    """Whitespace around the value is stripped -- only the digits are compared -- so a shell
    profile's trailing newline or a stray space in an env file does not accidentally disable it."""
    monkeypatch.setenv("AIPPT_LIVE_MODE", " 1 ")
    assert client.get("/api/context").json()["mode"] == "live"


def test_live_generate_injects_the_pin_and_returns_the_model_deck(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    _stub_worker(monkeypatch, returns=_live_deck())

    t0 = time.perf_counter()
    r = client.post("/api/generate", json={"prompt": "실적을 한 장으로 요약해줘"})
    elapsed = time.perf_counter() - t0

    assert r.status_code == 200, r.text
    ir = r.json()
    assert ir["title"] == "라이브 데크"
    assert [e["id"] for e in ir["elements"]] == ["box1"]

    pin = load_pin()
    assert ir["schema"] == "canvas-ir/1"
    assert ir["canvas"] == {"w": 960, "h": 540, "unit": "pt"}
    assert ir["target"]["template_sha256"] == pin["renderer_template"]["sha256"]
    assert ir["target"]["layout_index"] == 3

    validate(ir)  # the real validator -- the live deck is exportable
    assert elapsed < 1.0, "the live path slept the mock's 1.5s latency; it must not"


def test_the_live_deck_EXPORTS(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    _stub_worker(monkeypatch, returns=_live_deck())
    ir = client.post("/api/generate", json={"prompt": "x"}).json()
    r = client.post("/api/export", json=ir)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == PPTX_MIME
    assert len(Presentation(BytesIO(r.content)).slides) == 1


def test_live_generate_reports_a_worker_failure_as_502(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    _stub_worker(monkeypatch, raises=ai_adapter.CliWorkerError("quota exceeded"))
    r = client.post("/api/generate", json={"prompt": "x"})
    assert r.status_code == 502
    assert r.json()["error"] == "generate"
    assert "quota exceeded" in r.json()["detail"]


def test_live_generate_reports_a_schema_mismatch_as_502(client, monkeypatch):
    """CliWorkerSchemaError IS-A CliWorkerError -- 'the CLI answered but got the shape wrong' is
    still an upstream (502) failure from app.py's point of view, not a user error. (Bad-JSON/
    fence-wrapped output is ai_adapter.py's _extract_json_object()'s problem now, not app.py's --
    complete_json() never returns raw text, only a validated dict or one of these exceptions.)"""
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    _stub_worker(
        monkeypatch,
        raises=ai_adapter.CliWorkerSchemaError("no JSON object found in CLI output"),
    )
    r = client.post("/api/generate", json={"prompt": "x"})
    assert r.status_code == 502
    assert r.json()["error"] == "generate"


def test_live_generate_refuses_a_model_deck_on_the_logo(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    bad = json.loads(json.dumps(_LIVE_DECK))
    bad["elements"][0]["y"] = 10.0  # into HEADER_RESERVED (0..92)
    _stub_worker(monkeypatch, returns=bad)
    r = client.post("/api/generate", json={"prompt": "x"})
    assert r.status_code == 502
    assert r.json()["error"] == "model"
    assert "HEADER_RESERVED" in r.json()["detail"]


def test_live_generate_still_refuses_an_empty_prompt_immediately(client, monkeypatch):
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    called = {"n": 0}

    def must_not_call(self, *, system, user, schema, max_tokens=8000):
        called["n"] += 1
        return _live_deck()
    monkeypatch.setattr(ai_adapter.CliWorkerClient, "complete_json", must_not_call)
    r = client.post("/api/generate", json={"prompt": "   "})
    assert r.status_code == 422
    assert called["n"] == 0, "the model was called for an empty prompt"


def test_live_generate_uses_the_requested_profiles_backend_and_model(client, monkeypatch):
    """The profile picked on the request must reach CliWorkerClient's constructor, not just its
    thinking-level metadata -- this is the whole point of MODEL_PROFILES existing."""
    monkeypatch.setenv("AIPPT_LIVE_MODE", "1")
    monkeypatch.setenv("AIPPT_TRACK", "b")
    seen = {}

    def fake_init(self, *, track="b", backend="claude", model="sonnet", worker_role=None,
                  timeout_s=300):
        seen["track"] = track
        seen["backend"] = backend
        seen["model"] = model
        self.track, self.backend, self.model = track, backend, model

    def fake_complete_json(self, *, system, user, schema, max_tokens=8000):
        return _live_deck()

    monkeypatch.setattr(ai_adapter.CliWorkerClient, "__init__", fake_init)
    monkeypatch.setattr(ai_adapter.CliWorkerClient, "complete_json", fake_complete_json)

    r = client.post("/api/generate", json={"prompt": "x", "model": "gpt-5.6-sol"})
    assert r.status_code == 200, r.text
    assert seen == {"track": "b", "backend": "codex", "model": "gpt-5.6-sol"}


def test_without_the_env_var_generate_is_still_the_mock(client, monkeypatch):
    monkeypatch.delenv("AIPPT_LIVE_MODE", raising=False)
    a = client.post("/api/generate", json={"prompt": "a"}).json()
    b = client.post("/api/generate", json={"prompt": "완전히 다른 요청"}).json()
    assert a == b == mock_ir(load_pin())


def test_export_html_returns_standalone_html(client: TestClient):
    ir = load("minimal")
    r = client.post("/api/export_html", json=ir)
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]
    assert '<!DOCTYPE html>' in r.text
    assert 'class="slide-deck"' in r.text
    assert 'attachment; filename="deck.html"' in r.headers["content-disposition"]


def test_designs_api_save_load_delete(client: TestClient):
    name = "Test_Design_Phase1"
    ir = load("minimal")

    # Clean up before
    client.delete(f"/api/designs/{name}")

    # 1. List
    r_list = client.get("/api/designs")
    assert r_list.status_code == 200
    assert name not in r_list.json().get("designs", [])

    # 2. Save
    r_save = client.post("/api/designs", json={"name": name, "ir": ir})
    assert r_save.status_code == 200
    assert r_save.json()["name"] == name

    # 3. List after save
    r_list_after = client.get("/api/designs")
    assert r_list_after.status_code == 200
    assert name in r_list_after.json().get("designs", [])

    # 4. Load
    r_load = client.get(f"/api/designs/{name}")
    assert r_load.status_code == 200
    assert r_load.json()["schema"] == ir["schema"]

    # 5. Delete
    r_del = client.delete(f"/api/designs/{name}")
    assert r_del.status_code == 204

    # 6. Load deleted -> 404
    r_load_404 = client.get(f"/api/designs/{name}")
    assert r_load_404.status_code == 404

