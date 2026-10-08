#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The round-trip assertion suite (plan Sec 7.3). Phase 1A, task 1.7.

Every fixture is rendered, SAVED, and RE-OPENED before a single assertion runs. That is the whole
point: a shape that is correct in the in-memory tree and wrong on disk is precisely the failure mode
this suite exists to catch. The plan's own arrowhead snippet is an example -- it emits <a:tailEnd>
before <a:headEnd>, which is valid-looking Python, a well-formed lxml tree, and a file PowerPoint
offers to repair.

The geometry tolerance is 0 EMU, not "close". emu(u) = round(u * 12700) is exact by construction, so
any drift at all means the write path changed the number -- and a write path that rounds is a write
path that will keep rounding. The expected values below are therefore written as `round(u * 12700)`
longhand rather than by calling emu(): a test that computes its expectation with the code under test
proves only that the code agrees with itself.

Assertions are scoped to RENDERER-CREATED shapes. The title placeholder is inherited chrome -- it is
allowed to reference whatever the template says, and it is excluded from the accent1/shadow rules
(but not from the font rules, which apply to every run we write).

Colours and line widths are read straight out of the XML rather than through LineFormat.color, whose
getter documents a side-effect: it sets the line fill type to SOLID. Asserting through it would
mutate the very thing under test.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt

from aippt.canvas.geometry import (
    FOOTER_RESERVED,
    HEADER_RESERVED,
    element_bbox,
    intersects,
)
from aippt.canvas.pin import TITLE_PH_IDX, load_pin
from aippt.canvas.render_canvas import main, render
from aippt.canvas.style import ALIGN, DEFAULT_ALIGN, DEFAULT_VALIGN, EA, LATIN, VALIGN
from aippt.canvas.validate import StalePayloadError

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "canvas"

VALID = ["minimal", "shapes", "connectors", "rotated", "text", "edges"]

EMU_PER_UNIT = 12700  # stated longhand -- see the module docstring
FORBIDDEN = "4472C4"


def load(name: str) -> dict:
    return json.loads((FIXTURES / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def decks(tmp_path_factory):
    """Render each valid fixture once, save it, and hand back the RE-OPENED package."""
    out_dir = tmp_path_factory.mktemp("decks")
    built = {}
    for name in VALID:
        ir = load(name)
        prs = render(ir)
        path = out_dir / f"{name}.pptx"
        prs.save(str(path))
        built[name] = (ir, Presentation(str(path)))
    return built


def slide_of(decks, name):
    ir, prs = decks[name]
    assert len(prs.slides) == 1, f"{name}: expected exactly 1 slide"
    return ir, prs.slides[0]


def drawn(decks, name):
    """(element, shape) for every element the RENDERER created, found by name -- never by position."""
    ir, slide = slide_of(decks, name)
    shapes = {sh.name: sh for sh in slide.shapes}
    pairs = []
    for el in ir["elements"]:
        assert el["id"] in shapes, f"{name}: shape {el['id']!r} is missing from the export"
        pairs.append((el, shapes[el["id"]]))
    return pairs


def spPr(shape):
    return shape._element.spPr


def ln_of(shape):
    return spPr(shape).find(qn("a:ln"))


def srgb(el):
    """The `val` of the <a:srgbClr> under a fill element, or None."""
    if el is None:
        return None
    clr = el.find(qn("a:srgbClr"))
    return None if clr is None else clr.get("val")


def expected_pptx_runs(text):
    """The DrawingML runs an IR text object should produce, paragraph by paragraph.

    An IR run whose `t` contains '\\n' becomes SEVERAL <a:r> elements separated by <a:br/>, each
    carrying the original run's formatting. python-pptx's `paragraph.runs` returns only the <a:r>
    elements, so this expansion is what makes the two line up.
    """
    out = []
    for para in text["paragraphs"]:
        pieces = []
        for run in para["runs"]:
            for piece in run["t"].split("\n"):
                pieces.append({**run, "t": piece})
        out.append(pieces)
    return out


# --- the package -------------------------------------------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_exactly_one_slide(decks, name):
    _, prs = decks[name]
    assert len(prs.slides) == 1


@pytest.mark.parametrize("name", VALID)
def test_the_slide_is_on_layout_4(decks, name):
    """Adopting the wrong layout is worse than crashing: it looks fine, ships, and is discovered by
    the audience."""
    _, slide = slide_of(decks, name)
    partname = load_pin()["layout_signature"]["layout_partname"]
    assert str(slide.slide_layout.part.partname) == partname == "/ppt/slideLayouts/slideLayout4.xml"


@pytest.mark.parametrize("name", VALID)
def test_array_order_is_z_order(decks, name):
    """elements[0] is drawn first and therefore sits at the BACK. There is no `z` field, so the only
    thing carrying z-order is the order the shapes were appended in."""
    ir, slide = slide_of(decks, name)
    ids = [el["id"] for el in ir["elements"]]
    drawn_in_order = [sh.name for sh in slide.shapes if sh.name in set(ids)]
    assert drawn_in_order == ids


# --- geometry: 0 EMU ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_shape_geometry_is_exact_to_the_emu(decks, name):
    for el, shp in drawn(decks, name):
        if el["type"] == "connector":
            continue
        assert int(shp.left) == round(el["x"] * EMU_PER_UNIT), f'{el["id"]}.left'
        assert int(shp.top) == round(el["y"] * EMU_PER_UNIT), f'{el["id"]}.top'
        assert int(shp.width) == round(el["w"] * EMU_PER_UNIT), f'{el["id"]}.width'
        assert int(shp.height) == round(el["h"] * EMU_PER_UNIT), f'{el["id"]}.height'


@pytest.mark.parametrize("name", VALID)
def test_connector_endpoints_are_exact_to_the_emu(decks, name):
    """Endpoints, not the bounding box: a connector's off/ext plus flipH/flipV encode the direction,
    and it is the ENDPOINTS the author placed."""
    for el, cxn in drawn(decks, name):
        if el["type"] != "connector":
            continue
        assert int(cxn.begin_x) == round(el["x1"] * EMU_PER_UNIT), f'{el["id"]}.x1'
        assert int(cxn.begin_y) == round(el["y1"] * EMU_PER_UNIT), f'{el["id"]}.y1'
        assert int(cxn.end_x) == round(el["x2"] * EMU_PER_UNIT), f'{el["id"]}.x2'
        assert int(cxn.end_y) == round(el["y2"] * EMU_PER_UNIT), f'{el["id"]}.y2'


@pytest.mark.parametrize("name", VALID)
def test_rotation_survives(decks, name):
    for el, shp in drawn(decks, name):
        if el["type"] == "connector":
            continue
        assert shp.rotation == pytest.approx(float(el.get("rotation", 0)) % 360.0)


# --- fill, line, shadow, and the accent1 invariant ---------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_fill_is_what_the_ir_asked_for(decks, name):
    for el, shp in drawn(decks, name):
        if el["type"] == "connector":
            continue
        fill = el.get("fill")
        if fill is None:
            assert spPr(shp).find(qn("a:noFill")) is not None, f'{el["id"]}: expected <a:noFill/>'
        else:
            got = srgb(spPr(shp).find(qn("a:solidFill")))
            assert got == fill.lstrip("#").upper(), f'{el["id"]}: fill'


@pytest.mark.parametrize("name", VALID)
def test_line_is_what_the_ir_asked_for(decks, name):
    for el, shp in drawn(decks, name):
        if el["type"] == "connector":
            continue
        ln = ln_of(shp)
        assert ln is not None, f'{el["id"]}: <a:ln> must always be explicit, never inherited'

        stroke = el.get("stroke")
        if stroke is None:
            # "No line" must be an explicit <a:noFill/>. Width 0 would be a HAIRLINE, not no line.
            assert ln.find(qn("a:noFill")) is not None, f'{el["id"]}: expected an explicit no-line'
        else:
            assert srgb(ln.find(qn("a:solidFill"))) == stroke["color"].lstrip("#").upper()
            assert int(shp.line.width) == int(Pt(stroke["width"])), f'{el["id"]}: stroke width'


@pytest.mark.parametrize("name", VALID)
def test_connector_lines_are_what_the_ir_asked_for(decks, name):
    for el, cxn in drawn(decks, name):
        if el["type"] != "connector":
            continue
        ln = ln_of(cxn)
        assert ln is not None
        assert srgb(ln.find(qn("a:solidFill"))) == el["color"].lstrip("#").upper()
        assert int(cxn.line.width) == int(Pt(el["width"]))


@pytest.mark.parametrize("name", VALID)
def test_no_renderer_created_shape_inherits_the_office_drop_shadow(decks, name):
    """Override 3, the one everyone forgets. effectRef -> accent1 survives every amount of fill and
    line overriding, and there is no other shadow API. Connectors get it too."""
    for el, shp in drawn(decks, name):
        assert shp.shadow.inherit is False, f'{el["id"]}: still inheriting the shadow'


@pytest.mark.parametrize("name", VALID)
def test_the_style_block_is_gone(decks, name):
    """Override 5. <p:style> is optional in CT_Shape and in CT_Connector, and both are born with one
    pointing at schemeClr accent1."""
    for el, shp in drawn(decks, name):
        assert shp._element.find(qn("p:style")) is None, f'{el["id"]}: <p:style> survived'


@pytest.mark.parametrize("name", VALID)
def test_accent1_reaches_no_renderer_created_shape(decks, name):
    """Plan Sec 7.3's invariant, and the reason override 5 exists. Overrides 1-3 win at RENDER time
    but do not DELETE <p:style>, so without the strip this assertion could never pass on a deck that
    is in fact rendering correctly."""
    for el, shp in drawn(decks, name):
        xml = shp._element.xml
        assert "accent1" not in xml, f'{el["id"]}: accent1 in the shape XML'
        assert FORBIDDEN not in xml.upper(), f'{el["id"]}: {FORBIDDEN} in the shape XML'


# --- fonts: the latin/EA split -----------------------------------------------------------------


def rprs_of(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            rPr = run._r.find(qn("a:rPr"))
            assert rPr is not None
            yield rPr


@pytest.mark.parametrize("name", VALID)
def test_every_run_carries_both_typefaces(decks, name):
    """Century Gothic ships NO Hangul glyphs, so <a:ea> is not optional -- without it every Korean
    character falls back to whatever PowerPoint picks and the line width moves with it."""
    for el, shp in drawn(decks, name):
        if not shp.has_text_frame:
            continue
        for rPr in rprs_of(shp):
            latin = rPr.find(qn("a:latin"))
            ea = rPr.find(qn("a:ea"))
            assert latin is not None and latin.get("typeface") == LATIN, el["id"]
            assert ea is not None and ea.get("typeface") == EA, el["id"]


@pytest.mark.parametrize("name", VALID)
def test_the_ea_element_sits_in_schema_order(decks, name):
    """<a:ea> must FOLLOW <a:latin> and precede <a:cs>/<a:sym>/<a:hlink*>/<a:rtl>/<a:extLst>.
    insert_element_before APPENDS when it finds none of its successors, so an incomplete successor
    list silently puts <a:ea> in the wrong place and PowerPoint calls the file corrupt."""
    for el, shp in drawn(decks, name):
        if not shp.has_text_frame:
            continue
        for rPr in rprs_of(shp):
            children = list(rPr)
            latin = rPr.find(qn("a:latin"))
            ea = rPr.find(qn("a:ea"))
            assert children.index(latin) < children.index(ea), el["id"]


def test_the_title_run_carries_both_typefaces_too(decks):
    _, slide = slide_of(decks, "text")
    run = slide.placeholders[TITLE_PH_IDX].text_frame.paragraphs[0].runs[0]
    rPr = run._r.find(qn("a:rPr"))
    assert rPr.find(qn("a:latin")).get("typeface") == LATIN
    assert rPr.find(qn("a:ea")).get("typeface") == EA


# --- text ---------------------------------------------------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_the_text_frame_margins_mirror_the_css(decks, name):
    """7.2 / 3.6 logical units, not 0. Leave PowerPoint's defaults while the browser div has
    padding:0 and the two wrap at a DIFFERENT width -- an off-by-one-line overflow that presents as
    a coordinate bug."""
    for el, shp in drawn(decks, name):
        if not shp.has_text_frame:
            continue
        tf = shp.text_frame
        assert tf.margin_left == Pt(7.2), el["id"]
        assert tf.margin_right == Pt(7.2), el["id"]
        assert tf.margin_top == Pt(3.6), el["id"]
        assert tf.margin_bottom == Pt(3.6), el["id"]
        assert tf.word_wrap is True, el["id"]
        assert tf.auto_size == MSO_AUTO_SIZE.NONE, el["id"]  # autofit is out of scope (Sec 6.6)


@pytest.mark.parametrize("name", VALID)
def test_alignment_and_anchor_come_from_the_ir(decks, name):
    for el, shp in drawn(decks, name):
        if not shp.has_text_frame:
            continue
        text = el.get("text")
        if not text:
            continue
        assert shp.text_frame.vertical_anchor == VALIGN[text.get("valign", DEFAULT_VALIGN)]
        for para in shp.text_frame.paragraphs:
            assert para.alignment == ALIGN[text.get("align", DEFAULT_ALIGN)]


@pytest.mark.parametrize("name", VALID)
def test_every_run_property_survives(decks, name):
    for el, shp in drawn(decks, name):
        if not shp.has_text_frame or not el.get("text"):
            continue
        expected = expected_pptx_runs(el["text"])
        paragraphs = shp.text_frame.paragraphs
        assert len(paragraphs) == len(expected), f'{el["id"]}: paragraph count'

        for para, exp_runs in zip(paragraphs, expected):
            assert len(para.runs) == len(exp_runs), f'{el["id"]}: run count'
            assert para.space_before == Pt(0)
            assert para.space_after == Pt(0)
            assert para.line_spacing == pytest.approx(1.2)  # == CSS line-height: 1.2

            for run, exp in zip(para.runs, exp_runs):
                assert run.text == exp["t"]
                assert run.font.size == Pt(exp["size"])
                assert run.font.bold is exp.get("bold", False)
                assert run.font.italic is exp.get("italic", False)
                assert run.font.color.rgb == RGBColor.from_string(exp["color"].lstrip("#"))


def test_a_newline_becomes_a_line_break_not_a_space(decks):
    """A raw '\\n' left inside <a:t> is only whitespace to DrawingML -- it renders as a SPACE. The
    split into <a:r><a:br/><a:r> is the only thing that makes the character mean what the author
    typed."""
    _, slide = slide_of(decks, "text")
    shp = {sh.name: sh for sh in slide.shapes}["rich"]
    para = shp.text_frame.paragraphs[2]  # "line one\nline two"

    tags = [child.tag for child in para._p if child.tag in (qn("a:r"), qn("a:br"))]
    assert tags == [qn("a:r"), qn("a:br"), qn("a:r")]
    assert [r.text for r in para.runs] == ["line one", "line two"]


def test_the_break_carries_the_runs_formatting_across_it(decks):
    _, slide = slide_of(decks, "text")
    shp = {sh.name: sh for sh in slide.shapes}["rich"]
    before, after = shp.text_frame.paragraphs[2].runs
    assert before.font.size == after.font.size == Pt(14)
    assert before.font.color.rgb == after.font.color.rgb == RGBColor.from_string("000000")


def test_nbsp_is_preserved_and_never_collapsed_to_a_space(decks):
    _, slide = slide_of(decks, "text")
    shp = {sh.name: sh for sh in slide.shapes}["rich"]
    assert chr(0xA0) in shp.text_frame.paragraphs[3].runs[0].text


def test_korean_survives_the_round_trip(decks):
    _, slide = slide_of(decks, "text")
    shp = {sh.name: sh for sh in slide.shapes}["rich"]
    assert shp.text_frame.paragraphs[1].runs[0].text == "개발 테스트벤치 설계안 Rev 2"


# --- the title ----------------------------------------------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_the_title_lands_in_placeholder_idx_10(decks, name):
    ir, slide = slide_of(decks, name)
    ph = slide.placeholders[TITLE_PH_IDX]  # keyed by IDX on a slide; KeyError, not IndexError
    assert ph.text_frame.text == ir["title"]


def test_every_title_run_property_is_stated_explicitly(decks):
    """Do NOT trust the inherited placeholder style to match the editor's default."""
    _, slide = slide_of(decks, "minimal")
    para = slide.placeholders[TITLE_PH_IDX].text_frame.paragraphs[0]
    assert para.alignment == ALIGN["LEFT"]
    run = para.runs[0]
    assert run.font.size == Pt(30)
    assert run.font.bold is False
    assert run.font.color.rgb == RGBColor.from_string("000000")


# --- connectors and arrowheads ------------------------------------------------------------------

ARROW_ENDS = {
    "none": (False, False),
    "start": (True, False),
    "end": (False, True),
    "both": (True, True),
}


@pytest.mark.parametrize("name", VALID)
def test_arrowheads_are_present_exactly_where_the_ir_asked(decks, name):
    for el, cxn in drawn(decks, name):
        if el["type"] != "connector":
            continue
        ln = ln_of(cxn)
        head = ln.find(qn("a:headEnd"))
        tail = ln.find(qn("a:tailEnd"))
        want_head, want_tail = ARROW_ENDS[el["arrow"]]

        assert (head is not None) is want_head, f'{el["id"]}: <a:headEnd>'
        assert (tail is not None) is want_tail, f'{el["id"]}: <a:tailEnd>'
        for end in (head, tail):
            if end is not None:
                assert end.get("type") == "triangle"


def test_head_end_precedes_tail_end_in_the_serialised_xml(decks):
    """CORRECTION 1, and the reason this suite re-opens the saved file rather than inspecting the
    in-memory tree. CT_LineProperties._tag_seq is

        ... round bevel miter headEnd tailEnd extLst

    so <a:headEnd> MUST precede <a:tailEnd>. The plan's snippet appends tailEnd first, which for
    arrow:"both" emits them in reverse schema order -- and PowerPoint answers with a repair dialog.
    A fixture with arrow:"both" that survives a save/load is the proof.
    """
    both = [(el, cxn) for el, cxn in drawn(decks, "connectors") if el.get("arrow") == "both"]
    assert both, "the connectors fixture must exercise arrow:'both'"

    for el, cxn in both:
        ln = ln_of(cxn)
        children = list(ln)
        head = ln.find(qn("a:headEnd"))
        tail = ln.find(qn("a:tailEnd"))
        assert children.index(head) < children.index(tail), f'{el["id"]}: schema order violated'


def test_every_kind_and_arrow_combination_is_covered(decks):
    ir, _ = decks["connectors"]
    seen = {(el["kind"], el["arrow"]) for el in ir["elements"]}
    assert seen == {
        (kind, arrow)
        for kind in ("straight", "elbow")
        for arrow in ("none", "start", "end", "both")
    }


# --- the reserved regions, measured on the OUTPUT ------------------------------------------------


@pytest.mark.parametrize("name", VALID)
def test_no_exported_shape_intersects_a_reserved_region(decks, name):
    """Re-derived from the EXPORTED geometry, not from the input IR. Checking the input would only
    prove that validate() agrees with itself; the question here is where the shapes actually landed.
    """
    for el, shp in drawn(decks, name):
        if el["type"] == "connector":
            out = {
                "type": "connector",
                "x1": int(shp.begin_x) / EMU_PER_UNIT,
                "y1": int(shp.begin_y) / EMU_PER_UNIT,
                "x2": int(shp.end_x) / EMU_PER_UNIT,
                "y2": int(shp.end_y) / EMU_PER_UNIT,
            }
        else:
            out = {
                "type": el["type"],
                "x": int(shp.left) / EMU_PER_UNIT,
                "y": int(shp.top) / EMU_PER_UNIT,
                "w": int(shp.width) / EMU_PER_UNIT,
                "h": int(shp.height) / EMU_PER_UNIT,
                "rotation": float(shp.rotation),
            }
        box = element_bbox(out)
        assert not intersects(box, HEADER_RESERVED), f'{el["id"]} landed in HEADER_RESERVED'
        assert not intersects(box, FOOTER_RESERVED), f'{el["id"]} landed in FOOTER_RESERVED'


# --- the pin ------------------------------------------------------------------------------------


def test_a_stale_template_sha_is_refused_at_render_time():
    """Valid IR, wrong template. It passes validate() -- which is pure and knows nothing about the
    pin -- and is caught by render(). It is never rescaled or re-fitted."""
    with pytest.raises(StalePayloadError):
        render(load("reject_sha"))


# --- the CLI ------------------------------------------------------------------------------------


def test_the_cli_renders_a_deck(tmp_path):
    out = tmp_path / "out.pptx"
    assert main([str(FIXTURES / "minimal.json"), "-o", str(out)]) == 0
    assert out.is_file()
    assert len(Presentation(str(out)).slides) == 1


def test_the_cli_exits_non_zero_and_prints_every_violation(tmp_path, capsys):
    out = tmp_path / "nope.pptx"
    assert main([str(FIXTURES / "reject_accent1.json"), "-o", str(out)]) == 1

    err = capsys.readouterr().err
    assert "4 semantic violations" in err
    for eid in ("bad_fill", "bad_stroke", "bad_run", "bad_connector"):
        assert eid in err
    assert not out.exists(), "a rejected deck must not leave a file behind"


def test_the_cli_refuses_a_stale_payload(tmp_path):
    assert main([str(FIXTURES / "reject_sha.json"), "-o", str(tmp_path / "nope.pptx")]) == 1


def test_the_cli_reports_a_missing_deck(tmp_path, capsys):
    assert main([str(tmp_path / "absent.json"), "-o", str(tmp_path / "nope.pptx")]) == 2
    assert "not found" in capsys.readouterr().err
