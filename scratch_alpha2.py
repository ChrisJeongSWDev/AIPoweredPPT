from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
import sys
import os
sys.path.insert(0, os.path.abspath('src'))
from aippt.canvas.style import apply_color

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
shape.fill.solid()

apply_color(shape.fill.fore_color, "#FF000080")

print(shape.fill.fore_color._color._xClr.xml)
prs.save("test_apply_color.pptx")
